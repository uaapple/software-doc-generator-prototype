#!/usr/bin/env python3
"""Fail-closed semantic validation used by the TCSD host, never by the Agent."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from run_tcsd_pipeline_stage import simulation_backfill_evidence
from validate_agent_coverage_repair import (
    BRIEF_SCHEMA,
    IR_SCHEMA,
    PROPOSAL_SCHEMA,
    VALIDATION_SCHEMA,
    build_brief,
    validate_proposal,
)
from validate_tcsd_workbook import load_interface_names, validate_workbook


REPORT_SCHEMA = "tcsd-host-semantic-validation/v1"
COVERAGE_SCHEMA = "tcsd-coverage-report/v1"
SIMULATION_SCHEMA = "tcsd-simulation-result/v1"
ENVIRONMENT_SCHEMA = "tcsd-environment-gate/v2"
PROBE_PLAN_SCHEMA = "simulink-ut-state-probe-plan/v1"
PROBE_RESULT_SCHEMA = "simulink-ut-logical-mcdc-probe/v2"
REPAIR_CANDIDATE_SCHEMA = "tcsd-repair-candidate-validation/v1"
SYNTHESIS_SCHEMA = "simulink-ut-tcsd-coverage-ir-synthesis/v1"


def read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"JSON root must be an object: {path}")
    return value


def resolve_workspace_path(root: Path, candidate: str, label: str) -> Path:
    path = (root / candidate).resolve() if not Path(candidate).is_absolute() else Path(candidate).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError(f"{label} escaped the task workspace")
    if not path.is_file() or path.stat().st_size <= 0:
        raise ValueError(f"{label} is missing or empty: {candidate}")
    return path


def artifact_paths(request: dict[str, Any]) -> list[tuple[dict[str, Any], Path]]:
    root = Path(request["workspaceDir"]).resolve()
    return [
        (artifact, resolve_workspace_path(root, str(artifact.get("path") or ""), "artifact"))
        for artifact in request.get("artifacts", [])
        if isinstance(artifact, dict)
    ]


def json_artifacts(request: dict[str, Any]) -> list[tuple[dict[str, Any], Path, dict[str, Any]]]:
    values = []
    for artifact, path in artifact_paths(request):
        if artifact.get("kind") == "json":
            values.append((artifact, path, read_json(path)))
    return values


def find_json_schema(request: dict[str, Any], schema: str) -> tuple[Path, dict[str, Any]]:
    for _, path, value in json_artifacts(request):
        if value.get("schema") == schema:
            return path, value
    raise ValueError(f"required semantic artifact schema is missing: {schema}")


def find_workbook(request: dict[str, Any]) -> Path:
    for artifact, path in artifact_paths(request):
        if artifact.get("kind") == "xlsx":
            return path
    raise ValueError("semantic validation requires an XLSX workbook artifact")


def canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def validate_environment(request: dict[str, Any]) -> dict[str, Any]:
    _, gate = find_json_schema(request, ENVIRONMENT_SCHEMA)
    dependencies = gate.get("pythonDependencies")
    workspace_io = gate.get("workspaceIo")
    matlab = gate.get("matlab")
    simulink = gate.get("simulink")
    satk = gate.get("satkMcp")
    server = satk.get("server") if isinstance(satk, dict) else None
    modules = dependencies.get("modules") if isinstance(dependencies, dict) else None
    required_modules = {"yaml", "openpyxl"}
    if (
        gate.get("passed") is not True
        or not isinstance(modules, dict)
        or not required_modules.issubset(modules)
        or any(not str(modules[name].get("version") or "") for name in required_modules)
        or workspace_io != {
            "passed": True,
            "created": True,
            "readMatched": True,
            "deleted": True,
        }
        or not isinstance(matlab, dict)
        or matlab.get("passed") is not True
        or not matlab.get("nonce")
        or matlab.get("nonce") != gate.get("nonce")
        or not matlab.get("version")
        or not isinstance(simulink, dict)
        or simulink.get("passed") is not True
        or simulink.get("licenseAvailable") is not True
        or simulink.get("loaded") is not True
        or not simulink.get("version")
        or not isinstance(satk, dict)
        or satk.get("passed") is not True
        or satk.get("sentinelWritten") is not True
        or satk.get("nonceMatched") is not True
        or not satk.get("runner")
        or not isinstance(server, dict)
        or not Path(str(server.get("path") or "")).is_file()
        or not re.fullmatch(r"[a-f0-9]{64}", str(server.get("sha256") or ""))
        or int(server.get("sizeBytes") or 0) <= 0
    ):
        raise ValueError("environment canary evidence is incomplete or internally inconsistent")
    server_path = Path(str(server["path"])).resolve()
    if hashlib.sha256(server_path.read_bytes()).hexdigest() != server["sha256"]:
        raise ValueError("environment canary MCP server hash does not match the selected executable")
    return {
        "environmentSchema": ENVIRONMENT_SCHEMA,
        "dependencyModules": sorted(required_modules),
        "workspaceIoPassed": True,
        "matlabNonceSha256": hashlib.sha256(str(gate["nonce"]).encode()).hexdigest(),
        "simulinkLoaded": True,
        "satkSentinelWritten": True,
        "satkServerPath": str(server_path),
        "satkServerSha256": server["sha256"],
    }


def probe_reports(payload: dict[str, Any]) -> list[dict[str, Any]]:
    if payload.get("schema") == PROBE_RESULT_SCHEMA:
        return [payload]
    return [
        item
        for item in payload.values()
        if isinstance(item, dict) and item.get("schema") == PROBE_RESULT_SCHEMA
    ]


def validate_probe(request: dict[str, Any]) -> dict[str, Any]:
    _, plan = find_json_schema(request, PROBE_PLAN_SCHEMA)
    tests = plan.get("tests")
    summary = plan.get("summary")
    if not isinstance(tests, list) or not isinstance(summary, dict):
        raise ValueError("state Probe plan does not expose tests and summary")
    candidate_count = int(summary.get("candidate_count") or 0)
    if candidate_count != len(tests):
        raise ValueError("state Probe candidate count does not match plan tests")
    planned: dict[str, set[int]] = {}
    for test in tests:
        test_id = str(test.get("test_id") or "")
        steps = test.get("steps")
        indices = {
            int(step.get("index") or 0)
            for step in steps
            if isinstance(step, dict) and int(step.get("index") or 0) > 0
        } if isinstance(steps, list) else set()
        if not test_id or not indices or test_id in planned:
            raise ValueError("state Probe plan contains an invalid candidate")
        planned[test_id] = indices
    evidence = request.get("evidence") if isinstance(request.get("evidence"), dict) else {}
    if int(evidence.get("candidateCount") or 0) != candidate_count:
        raise ValueError("Agent candidateCount does not match the parsed Probe plan")
    if candidate_count == 0:
        if evidence.get("probeExecuted") is not False:
            raise ValueError("empty Probe plan must explicitly report probeExecuted=false")
        return {"candidateCount": 0, "probeExecuted": False, "observationCount": 0}
    if evidence.get("probeExecuted") is not True:
        raise ValueError("Probe candidates exist but probeExecuted is not true")
    reports: list[dict[str, Any]] = []
    for _, _, payload in json_artifacts(request):
        reports.extend(probe_reports(payload))
    if not reports:
        raise ValueError("Probe candidates exist but no actual Probe result is present")
    observed: dict[str, set[int]] = {}
    observation_count = 0
    for report in reports:
        observations = report.get("observations")
        if not isinstance(observations, list):
            raise ValueError("Probe result observations must be an array")
        for observation in observations:
            if not isinstance(observation, dict):
                continue
            test_id = str(observation.get("test_id") or "")
            step_index = int(observation.get("step_index") or 0)
            vectors = observation.get("vectors")
            valid_vectors = [
                vector
                for vector in vectors.values()
                if isinstance(vectors, dict)
                and isinstance(vector, dict)
                and vector.get("ok") is True
                and isinstance(vector.get("values"), list)
            ] if isinstance(vectors, dict) else []
            mps_blocked = observation.get("prediction_status") == "simulation_error_mps_selector"
            if (
                test_id not in planned
                or step_index not in planned[test_id]
                or not isinstance(observation.get("inputs"), dict)
                or (not valid_vectors and not mps_blocked)
                or observation.get("prediction_status") in {"target_unavailable", "simulation_mismatch"}
            ):
                raise ValueError("Probe observation is missing executed values or contradicts its planned target")
            observed.setdefault(test_id, set()).add(step_index)
            observation_count += 1
    for test_id, indices in planned.items():
        if observed.get(test_id) != indices:
            raise ValueError(f"Probe candidate {test_id} was not observed for every planned step")
    return {
        "candidateCount": candidate_count,
        "probeExecuted": True,
        "observationCount": observation_count,
    }


def count_action_steps(workbook: Path) -> int:
    wb = load_workbook(workbook, read_only=True, data_only=False)
    try:
        if "TCSD" not in wb.sheetnames:
            return 0
        ws = wb["TCSD"]
        count = 0
        for row in range(1, ws.max_row + 1):
            if str(ws.cell(row, 3).value or "").strip() != "Test":
                continue
            count += sum(
                1
                for line in str(ws.cell(row, 7).value or "").splitlines()
                if line.strip().startswith("[+")
            )
        return count
    finally:
        wb.close()


def validate_workbook_stage(request: dict[str, Any], require_exp_values: bool) -> dict[str, Any]:
    workbook = find_workbook(request)
    interface_path = resolve_workspace_path(
        Path(request["workspaceDir"]),
        str(request.get("interfacePath") or ""),
        "model interface",
    )
    root_inputs, root_outputs = load_interface_names(str(interface_path))
    report = validate_workbook(
        workbook,
        root_inputs,
        root_outputs,
        require_exp_values=require_exp_values,
    )
    template_path = Path(str(request.get("templatePath") or "")).resolve()
    if template_path.is_file() and hashlib.sha256(workbook.read_bytes()).digest() == hashlib.sha256(template_path.read_bytes()).digest():
        raise ValueError("unfilled bundled TCSD template is not a generated workbook")
    action_step_count = count_action_steps(workbook)
    if (
        report.get("status") != "passed"
        or int(report.get("test_count") or 0) < 1
        or action_step_count < int(report.get("test_count") or 0)
    ):
        first = (report.get("errors") or [{"code": "missing_executable_tests"}])[0]
        raise ValueError(f"TCSD workbook semantic validation failed: {first.get('code')}")
    return {
        "testCount": int(report["test_count"]),
        "inputAssignmentCount": int(report.get("input_assignment_count") or 0),
        "parameterAssignmentCount": int(report.get("parameter_assignment_count") or 0),
        "expValueCount": int(report.get("exp_value_count") or 0),
        "actionStepCount": action_step_count,
        "workbookSha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
    }


def validate_simulation(request: dict[str, Any]) -> dict[str, Any]:
    workbook = find_workbook(request)
    simulation_path, simulation = find_json_schema(request, SIMULATION_SCHEMA)
    canonical_evidence = simulation_backfill_evidence(simulation, workbook)
    claimed = request.get("evidence") if isinstance(request.get("evidence"), dict) else {}
    for key in ("simulationValueCount", "workbookBackfillCount", "expValueCount"):
        expected = canonical_evidence["workbookBackfillCount"] if key == "expValueCount" else canonical_evidence.get(key)
        if int(claimed.get(key) or 0) != int(expected or 0):
            raise ValueError(f"Agent simulation evidence does not match host-parsed {key}")
    if canonical(claimed.get("caseOutputCounts")) != canonical(canonical_evidence.get("caseOutputCounts")):
        raise ValueError("Agent simulation evidence does not match host-parsed caseOutputCounts")
    normalize_items = lambda items: [
        {
            "row": int(item.get("row") or 0),
            "testId": str(item.get("testId") or ""),
            "step": int(item.get("step") or 0),
            "output": str(item.get("output") or ""),
            "value": float(item.get("value")),
        }
        for item in items
        if isinstance(item, dict)
    ] if isinstance(items, list) else []
    if canonical(normalize_items(claimed.get("backfillItems"))) != canonical(
        normalize_items(canonical_evidence.get("backfillItems"))
    ):
        raise ValueError("Agent simulation evidence does not match host-parsed backfillItems")
    if str(claimed.get("simulationResult") or "") != str(
        next(
            artifact.get("path")
            for artifact, path in artifact_paths(request)
            if path == simulation_path
        )
    ):
        raise ValueError("Agent simulationResult does not reference the parsed simulation artifact")
    return {
        **canonical_evidence,
        "expValueCount": canonical_evidence["workbookBackfillCount"],
        "simulationSha256": hashlib.sha256(simulation_path.read_bytes()).hexdigest(),
    }


def metric(record: dict[str, Any], name: str, threshold: float) -> dict[str, Any]:
    raw = record.get(name)
    if not isinstance(raw, dict):
        raise ValueError(f"coverage metric is missing: {name}")
    covered = float(raw.get("covered"))
    total = float(raw.get("total"))
    percent = float(raw.get("percent"))
    if not all(math.isfinite(value) for value in (covered, total, percent)):
        raise ValueError(f"coverage metric is non-finite: {name}")
    if covered < 0 or total < 0 or covered > total or percent < 0 or percent > 100:
        raise ValueError(f"coverage metric range is invalid: {name}")
    expected = 100.0 if total == 0 else 100.0 * covered / total
    if not math.isclose(percent, expected, rel_tol=1e-7, abs_tol=1e-7):
        raise ValueError(f"coverage percent is not derived from covered/total: {name}")
    passed = percent >= threshold
    if raw.get("passed") is not passed:
        raise ValueError(f"coverage passed flag contradicts percent: {name}")
    return {"covered": covered, "total": total, "percent": percent, "passed": passed}


def validate_coverage(request: dict[str, Any]) -> dict[str, Any]:
    path, report = find_json_schema(request, COVERAGE_SCHEMA)
    models = report.get("models")
    threshold = float(request.get("coverageThreshold") or 80)
    if not isinstance(models, dict) or not models:
        raise ValueError("coverage report has no model records")
    normalized: dict[str, Any] = {}
    positive_totals = 0
    for model, record in models.items():
        if not model or not isinstance(record, dict):
            raise ValueError("coverage report contains an invalid model record")
        normalized[str(model)] = {
            name: metric(record, name, threshold)
            for name in ("condition", "decision", "mcdc")
        }
        positive_totals += sum(
            1
            for value in normalized[str(model)].values()
            if value["total"] > 0
        )
    if positive_totals < 1:
        raise ValueError("coverage report contains no measured obligations")
    return {
        "coverage": {"models": normalized},
        "coverageReportPath": str(
            next(artifact.get("path") for artifact, candidate in artifact_paths(request) if candidate == path)
        ),
        "coverageReportSha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def validate_repair(request: dict[str, Any]) -> dict[str, Any]:
    _, brief = find_json_schema(request, BRIEF_SCHEMA)
    _, proposal = find_json_schema(request, PROPOSAL_SCHEMA)
    _, validation = find_json_schema(request, VALIDATION_SCHEMA)
    _, repair_ir = find_json_schema(request, IR_SCHEMA)
    root = Path(request["workspaceDir"])
    evidence = brief.get("evidence") if isinstance(brief.get("evidence"), dict) else {}
    coverage_ref = str(evidence.get("coverageReport") or "")
    traces_ref = str(evidence.get("logicalTraces") or "")
    coverage_ir_ref = str(evidence.get("coverageIr") or "")
    interface_ref = str(evidence.get("modelInterface") or "")
    coverage_path = resolve_workspace_path(root, coverage_ref, "stage 10 source coverage report")
    traces_path = resolve_workspace_path(root, traces_ref, "stage 10 source logical traces")
    resolve_workspace_path(root, coverage_ir_ref, "stage 10 source Coverage IR")
    interface_path = resolve_workspace_path(root, interface_ref, "stage 10 source model interface")
    expected_brief = build_brief(
        job_id=str(request.get("jobId") or ""),
        model=str(brief.get("model") or ""),
        coverage=read_json(coverage_path),
        traces=read_json(traces_path),
        coverage_ir_path=coverage_ir_ref,
        coverage_report_path=coverage_ref,
        trace_path=traces_ref,
        interface_path=interface_ref,
        threshold=float(request.get("coverageThreshold") or 80),
    )
    if canonical(brief) != canonical(expected_brief):
        raise ValueError("Agent repair brief does not match host-rebuilt measured coverage deficits")
    interface_path = resolve_workspace_path(
        root,
        str(request.get("interfacePath") or ""),
        "model interface",
    )
    if interface_path != resolve_workspace_path(root, interface_ref, "stage 10 brief model interface"):
        raise ValueError("stage 10 repair brief references a different model interface")
    expected_ir, expected_validation = validate_proposal(proposal, brief, read_json(interface_path))
    if canonical(repair_ir) != canonical(expected_ir):
        raise ValueError("Agent repair Coverage IR does not match independent proposal validation")
    if canonical(validation) != canonical(expected_validation):
        raise ValueError("Agent repair validation report does not match independent host validation")

    synthesis: dict[str, Any] | None = None
    candidate: dict[str, Any] | None = None
    for _, _, value in json_artifacts(request):
        if value.get("schema") == SYNTHESIS_SCHEMA:
            synthesis = value
        elif value.get("schema") == REPAIR_CANDIDATE_SCHEMA:
            candidate = value
    repair = request.get("repair") if isinstance(request.get("repair"), dict) else {}
    applied = repair.get("applied") is True
    accepted = int(validation.get("acceptedCandidateCount") or 0)
    unresolved = int(validation.get("unresolvedCount") or 0)
    added = int(synthesis.get("added") or 0) if synthesis else 0
    reason = str(repair.get("reason") or "")
    allowed_reasons = {
        "agent_targeted_candidates_validated_and_appended",
        "agent_reported_specific_unresolved_deficits",
        "agent_candidates_duplicate_existing_tests",
        "agent_candidate_simulation_failed",
    }
    if reason not in allowed_reasons:
        raise ValueError("stage 10 repair reason is not a specific validated outcome")
    if applied:
        if accepted < 1 or added < 1 or not isinstance(candidate, dict) or candidate.get("passed") is not True:
            raise ValueError("applied stage 10 repair lacks accepted, synthesized, and simulated candidates")
        workbook_details = validate_workbook_stage(request, require_exp_values=True)
        simulation_details = validate_simulation(request)
    else:
        workbook_details = {}
        simulation_details = {}
        if reason == "agent_reported_specific_unresolved_deficits" and (accepted != 0 or unresolved < 1):
            raise ValueError("unresolved stage 10 result lacks specific unresolved deficit evidence")
        if reason == "agent_candidates_duplicate_existing_tests" and (accepted < 1 or not synthesis or added != 0):
            raise ValueError("duplicate stage 10 result lacks deterministic deduplication evidence")
        if reason == "agent_candidate_simulation_failed" and (
            accepted < 1 or not isinstance(candidate, dict) or candidate.get("passed") is not False
        ):
            raise ValueError("failed stage 10 candidate lacks deterministic simulation failure evidence")
    return {
        "proposalItemCount": int(validation.get("proposalItemCount") or 0),
        "acceptedCandidateCount": accepted,
        "unresolvedCount": unresolved,
        "synthesisAddedCount": added,
        "candidateValidationPassed": bool(candidate and candidate.get("passed") is True),
        **workbook_details,
        **simulation_details,
    }


def validate(request: dict[str, Any]) -> dict[str, Any]:
    stage = int(request.get("stageIndex") or 0)
    details: dict[str, Any] = {}
    if stage == 2:
        details.update(validate_environment(request))
    elif stage == 6:
        details.update(validate_probe(request))
    elif stage == 7:
        details.update(validate_workbook_stage(request, require_exp_values=False))
    elif stage == 8:
        details.update(validate_workbook_stage(request, require_exp_values=True))
        details.update(validate_simulation(request))
    elif stage == 9:
        details.update(validate_coverage(request))
    elif stage == 10:
        details.update(validate_repair(request))
    elif stage == 11:
        details.update(validate_workbook_stage(request, require_exp_values=True))
        details.update(validate_simulation(request))
        details.update(validate_coverage(request))
    return {
        "schema": REPORT_SCHEMA,
        "stageIndex": stage,
        "passed": True,
        "details": details,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--request", required=True)
    args = parser.parse_args()
    try:
        request = read_json(Path(args.request))
        print(json.dumps(validate(request), ensure_ascii=False))
        return 0
    except Exception as exc:
        print(
            json.dumps(
                {
                    "schema": REPORT_SCHEMA,
                    "stageIndex": 0,
                    "passed": False,
                    "message": str(exc),
                },
                ensure_ascii=False,
            ),
            file=sys.stderr,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
