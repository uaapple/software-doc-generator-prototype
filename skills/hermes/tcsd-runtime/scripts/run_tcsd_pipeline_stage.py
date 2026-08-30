#!/usr/bin/env python3
"""Shared deterministic runtime invoked by one atomic TCSD stage skill."""
from __future__ import annotations

import argparse, hashlib, importlib.util, json, math, os, re, secrets, shutil, subprocess, sys
from pathlib import Path
from typing import Any

INPUT_SCHEMA = "tcsd-agent-stage-input/v1"
RESULT_SCHEMA = "tcsd-agent-stage-result/v1"
STAGE6_PROBE_TIMEOUT_BASE_SECONDS = 600
STAGE6_PROBE_TIMEOUT_PER_CANDIDATE_SECONDS = 5
STAGE6_PROBE_TIMEOUT_MAX_SECONDS = 3600
STAGE11_PROBE_TIMEOUT_BASE_SECONDS = 600
STAGE11_PROBE_TIMEOUT_PER_CASE_SECONDS = 30
STAGE11_PROBE_TIMEOUT_MAX_SECONDS = 3600

def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path); module = importlib.util.module_from_spec(spec); assert spec.loader; spec.loader.exec_module(module); return module

def read_json(path: Path) -> dict[str, Any]: return json.loads(path.read_text(encoding="utf-8"))
def write_json(path: Path, value: Any) -> None: path.parent.mkdir(parents=True, exist_ok=True); path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
def file_sha256(path: Path) -> str: return hashlib.sha256(path.read_bytes()).hexdigest()
def probe_timeout_seconds(count: int, *, base: int, per_unit: int, maximum: int) -> int:
    configured = os.environ.get("SATK_GATEWAY_TIMEOUT_SECONDS", "").strip()
    configured_floor = base
    if configured:
        try:
            configured_floor = max(1, int(float(configured)))
        except ValueError:
            configured_floor = base
    adaptive = base + max(0, count) * per_unit
    return min(maximum, max(configured_floor, adaptive))
def stage6_probe_timeout_seconds(candidate_count: int) -> int:
    return probe_timeout_seconds(candidate_count, base=STAGE6_PROBE_TIMEOUT_BASE_SECONDS, per_unit=STAGE6_PROBE_TIMEOUT_PER_CANDIDATE_SECONDS, maximum=STAGE6_PROBE_TIMEOUT_MAX_SECONDS)
def stage11_probe_timeout_seconds(case_count: int) -> int:
    return probe_timeout_seconds(case_count, base=STAGE11_PROBE_TIMEOUT_BASE_SECONDS, per_unit=STAGE11_PROBE_TIMEOUT_PER_CASE_SECONDS, maximum=STAGE11_PROBE_TIMEOUT_MAX_SECONDS)

PROBE_CAPABILITY_SCHEMA = "tcsd-probe-capability/v1"
PROBE_CAPABILITY_CACHE_SCHEMA = "tcsd-probe-capability-cache/v1"
PROBE_CAPABILITY_STRATEGY_VERSION = "v1"
PROBE_PRECHECK_TIMEOUT_SECONDS = 900


def workspace_model_fingerprint(model_path: Path, root: Path) -> str:
    """Cache key for the capability precheck: model + every workspace library +
    the precheck strategy itself. A library or model change invalidates prior
    verdicts (an operator probeable under one library revision may be denied
    under the next)."""
    import hashlib as _hashlib
    digest = _hashlib.sha256()
    candidates = [model_path]
    for slx in sorted(root.glob("*.slx")):
        if slx.resolve() != model_path.resolve():
            candidates.append(slx)
    for path in candidates:
        if path.is_file():
            digest.update(path.name.encode())
            digest.update(b"\0")
            digest.update(_hashlib.sha256(path.read_bytes()).digest())
            digest.update(b"\0")
    digest.update(PROBE_CAPABILITY_STRATEGY_VERSION.encode())
    precheck_script = scripts() / "probe_capability_precheck.m"
    if precheck_script.is_file():
        digest.update(_hashlib.sha256(precheck_script.read_bytes()).digest())
    return digest.hexdigest()


def probe_capability_cache_path(out: Path) -> Path:
    return out / ".tcsd-runtime" / "probe-capability-cache.json"


def load_probe_capability_cache(out: Path, model: str, fingerprint: str):
    try:
        payload = read_json(probe_capability_cache_path(out))
    except (OSError, ValueError):
        return None
    if not isinstance(payload, dict) or payload.get("schema") != PROBE_CAPABILITY_CACHE_SCHEMA:
        return None
    if payload.get("model") != model or payload.get("fingerprint") != fingerprint:
        return None
    if payload.get("strategyVersion") != PROBE_CAPABILITY_STRATEGY_VERSION:
        return None
    capabilities = payload.get("capabilities")
    return capabilities if isinstance(capabilities, dict) else None


def save_probe_capability_cache(out: Path, model: str, fingerprint: str, capabilities: dict) -> None:
    write_json(probe_capability_cache_path(out), {
        "schema": PROBE_CAPABILITY_CACHE_SCHEMA,
        "model": model,
        "fingerprint": fingerprint,
        "strategyVersion": PROBE_CAPABILITY_STRATEGY_VERSION,
        "capabilities": capabilities,
        "computedAt": __import__("datetime").datetime.now().isoformat(),
    })


def build_probe_capability_request(mapping: Path, request_path: Path) -> int:
    """Extract the operators to classify from the logical mapping report."""
    payload = read_json(mapping)
    items = []
    if isinstance(payload, dict):
        operators = payload.get("operators")
        if isinstance(operators, dict):
            operators = [operators]
        if isinstance(operators, list):
            items = [item for item in operators if isinstance(item, dict)]
    request = {
        "schema": "tcsd-probe-capability-request/v1",
        "model": payload.get("model") if isinstance(payload, dict) else "",
        "operators": [
            {"id": str(item.get("id") or item.get("sid") or item.get("block_path") or ""),
             "block_path": str(item.get("block_path") or "")}
            for item in items
        ],
    }
    write_json(request_path, request)
    return len(request["operators"])


def run_probe_capability_precheck(job: dict, model: str, root: Path, out: Path,
                                  mapping: Path, python=sys.executable):
    """Classify probe observability for every candidate operator before any
    simulation. Returns (capability_path_or_None, evidence dict)."""
    model_path = Path(job["input"]["modelSlxPath"])
    fingerprint = workspace_model_fingerprint(model_path, root)
    cached = load_probe_capability_cache(out, model, fingerprint)
    if cached is not None:
        return None, {"precheckExecuted": False, "cache": "hit",
                      "strategyVersion": PROBE_CAPABILITY_STRATEGY_VERSION,
                      "capabilities": cached}
    precheck_fixture = os.environ.get("TCSD_PIPELINE_PRECHECK_CAPABILITY_FIXTURE", "").strip()
    capability_output = out / f"{model}_probe_capability.json"
    if precheck_fixture:
        # Test/offline fixture: same pattern as TCSD_PIPELINE_PROBE_RESULTS_FIXTURE.
        shutil.copy2(precheck_fixture, capability_output)
        payload = read_json(capability_output)
        if payload.get("schema") != PROBE_CAPABILITY_SCHEMA:
            raise RuntimeError(f"probe capability fixture schema invalid: {payload.get('schema')!r}")
        capabilities = payload.get("capabilities") if isinstance(payload.get("capabilities"), dict) else {}
        return capability_output, {"precheckExecuted": True, "cache": "fixture",
                                   "strategyVersion": PROBE_CAPABILITY_STRATEGY_VERSION,
                                   "capabilities": capabilities}
    request_path = out / ".tcsd-runtime" / "probe-capability-request.json"
    operator_count = build_probe_capability_request(mapping, request_path)
    entry = quality.write_matlab_entry(out / f"{model}_probe_capability_entry.m", "\n".join([
        f"rootDir = '{str(root).replace(chr(39), chr(39) * 2)}';",
        f"addpath('{str(scripts()).replace(chr(39), chr(39) * 2)}');",
        f"probe_capability_precheck(rootDir, '{model}', "
        f"'{str(request_path).replace(chr(39), chr(39) * 2)}', "
        f"'{str(capability_output).replace(chr(39), chr(39) * 2)}');",
    ]))
    quality.run_satk(python, scripts(), entry, root_dir=root,
                     gateway_timeout_seconds=PROBE_PRECHECK_TIMEOUT_SECONDS)
    payload = read_json(capability_output)
    if payload.get("schema") != PROBE_CAPABILITY_SCHEMA:
        raise RuntimeError(f"probe capability precheck produced {payload.get('schema')!r}")
    if payload.get("strategyVersion") != PROBE_CAPABILITY_STRATEGY_VERSION:
        raise RuntimeError("probe capability precheck strategy version mismatch")
    capabilities = payload.get("capabilities") if isinstance(payload.get("capabilities"), dict) else {}
    save_probe_capability_cache(out, model, fingerprint, capabilities)
    unprobeable = sum(1 for item in capabilities.values()
                      if isinstance(item, dict) and item.get("strategy") == "unprobeable")
    evidence = {
        "precheckExecuted": True,
        "cache": "miss",
        "strategyVersion": PROBE_CAPABILITY_STRATEGY_VERSION,
        "operatorCount": operator_count,
        "unprobeableCount": unprobeable,
        "capabilities": capabilities,
    }
    return capability_output, evidence
def probe_observation_reconciliation(probe_results: Path) -> dict:
    """Terminal-state tally across every probe report (reconciliation contract
    tcsd-probe-recon/v1). gapCount>0 forces the stage to end as `partial` so
    the platform sees "completed with registered gaps" instead of a hard
    failure — the authoritative verdict remains the Stage 9 measured
    coverage."""
    payload = read_json(probe_results)
    recon = {
        "observedCount": 0, "mismatchCount": 0, "transientFailedCount": 0,
        "notExecutedCount": 0, "mpsBlockedCount": 0, "otherCount": 0, "gapCount": 0,
    }
    reports = []
    if isinstance(payload, dict):
        for value in payload.values():
            if isinstance(value, dict) and value.get("schema") == "simulink-ut-logical-mcdc-probe/v2":
                reports.append(value)
    for report in reports:
        observations = report.get("observations")
        if not isinstance(observations, list):
            continue
        for observation in observations:
            if not isinstance(observation, dict):
                continue
            status = str(observation.get("prediction_status") or "")
            if status in ("observed", "matched_prediction"):
                recon["observedCount"] += 1
            elif status == "simulation_mismatch":
                recon["mismatchCount"] += 1
            elif status == "transient_failed":
                recon["transientFailedCount"] += 1
            elif status == "not_executed_with_reason":
                recon["notExecutedCount"] += 1
            elif status == "simulation_error_mps_selector":
                recon["mpsBlockedCount"] += 1
            else:
                recon["otherCount"] += 1
    recon["gapCount"] = (recon["mismatchCount"] + recon["transientFailedCount"]
                         + recon["notExecutedCount"] + recon["mpsBlockedCount"]
                         + recon["otherCount"])
    return recon


def planning_mapping_assessment(raw: dict[str, Any], obligations: Path, root: Path) -> dict[str, Any]:
    assessment = dict(raw)
    raw_status = str(assessment.pop("status", "failed"))
    assessment.update({
        "schema": "tcsd-planning-mapping-assessment/v1",
        "authority": "planning",
        "status": "satisfied" if raw_status == "passed" else "advisory",
        "blocking": False,
        "assessment": "complete" if raw_status == "passed" else "gaps-observed",
        "sourceObligations": {
            "path": str(obligations.relative_to(root)),
            "sha256": file_sha256(obligations),
        },
        "supersededBy": {
            "stageIndex": 9,
            "authority": "measured-simulink-coverage",
            "reason": "Static assignment matching is planning evidence; actual Condition/Decision/MC/DC is authoritative.",
        },
    })
    return assessment
def artifact(root: Path, path: Path, kind: str = "json", role: str = "evidence") -> dict[str, Any]: return {"path": path.resolve().relative_to(root.resolve()).as_posix(), "kind": kind, "role": role}
def model_name(job: dict[str, Any]) -> str: return Path(job["input"]["modelSlxPath"]).stem
def outputs(job: dict[str, Any]) -> Path: return Path(job["input"]["outputDir"]).resolve()
def workspace(job: dict[str, Any]) -> Path: return Path(job["input"]["workspaceDir"]).resolve()
def scripts() -> Path: return Path(__file__).resolve().parent
def state_path(job: dict[str, Any]) -> Path: return outputs(job) / ".tcsd-runtime" / "runner-state.json"
def load_state(job: dict[str, Any]) -> dict[str, Any]: return read_json(state_path(job)) if state_path(job).exists() else {"schema": "tcsd-stage-runner-state/v1", "jobId": job["jobId"], "resources": job.get("resources", {})}
def save_state(job: dict[str, Any], state: dict[str, Any]) -> None: write_json(state_path(job), state)
def finish(job: dict[str, Any], stage: int, *, status="completed", summary="", artifacts=None, **extra):
    payload = {"schema": RESULT_SCHEMA, "jobId": job["jobId"], "stageIndex": stage, "status": status, "summary": summary, "artifacts": artifacts or [], **extra}
    write_json(Path(job["_stageResultPath"]), payload)

def ensure_within(root: Path, candidate: Path, label: str) -> Path:
    resolved = candidate.resolve()
    if not resolved.is_relative_to(root.resolve()): raise RuntimeError(f"{label} is outside the task workspace: {resolved}")
    return resolved

def hard_error_code(stage: int, error: BaseException) -> str:
    if isinstance(error, subprocess.TimeoutExpired): return "tcsd_stage_timeout"
    if stage == 1: return "tcsd_input_invalid"
    if stage == 2: return "tcsd_environment_gate_failed"
    return "tcsd_stage_runtime_failed"

PUBLIC_ERROR_LIMIT = 2000
SECRET_ASSIGNMENT_RE = re.compile(
    r"(?i)\b(api[_-]?key|authorization|bearer|password|secret|token)\b(\s*[:=]\s*|\s+)([^\s,;]+)"
)

def public_error_text(value: Any) -> str:
    text = str(value or "").replace("\x00", "").replace("\r", "\n")
    text = SECRET_ASSIGNMENT_RE.sub(lambda match: f"{match.group(1)}{match.group(2)}[REDACTED]", text)
    text = " ".join(text.split())
    return text[:PUBLIC_ERROR_LIMIT]

def mcp_error_text(output: str) -> str:
    try:
        payload = json.loads(output)
    except (json.JSONDecodeError, TypeError):
        return ""
    error = payload.get("error")
    if isinstance(error, dict) and error.get("message"):
        return public_error_text(error["message"])
    result = payload.get("result")
    if not isinstance(result, dict) or result.get("isError") is not True:
        return ""
    content = result.get("content")
    if not isinstance(content, list):
        return ""
    messages = [
        public_error_text(item.get("text"))
        for item in content
        if isinstance(item, dict) and item.get("type") == "text" and item.get("text")
    ]
    return " ".join(message for message in messages if message)

def run(command: list[str], cwd: Path) -> None:
    subprocess.run(command, cwd=cwd, check=True)

def run_satk(command: list[str], cwd: Path, *, stage: int, context: str) -> None:
    transport = os.environ.get("TCSD_GATEWAY_TRANSPORT", "").strip()
    if transport and len(command) >= 2 and Path(command[1]).name == "satk_eval.py":
        command = [transport, *command[1:]]
    try:
        subprocess.run(command, cwd=cwd, check=True, capture_output=True, text=True)
    except subprocess.CalledProcessError as error:
        detail = mcp_error_text(error.stdout or "")
        if not detail:
            detail = public_error_text(error.stderr or "")
        if not detail:
            detail = f"satk_eval exited with code {error.returncode}"
        raise RuntimeError(f"Stage {stage:02d} {context}: SATK/MCP failed: {detail}") from None
def matlab_root_path(inp: dict[str, Any]) -> Path:
    return Path(
        os.environ.get("MATLAB_ROOT")
        or os.environ.get("SATK_MATLAB_ROOT")
        or inp.get("matlabRoot", "C:/Program Files/MATLAB/R2026a")
    )
def matlab_cell(items: list[str]) -> str: return "{" + ",".join("'" + item.replace("'", "''") + "'" for item in items) + "}"
def matlab_string(value: str) -> str: return "'" + value.replace("'", "''") + "'"
def stage4_matlab_code(*, root: Path, scripts_dir: Path, interface: Path, model: str, mat_name: str, init_scripts: list[str]) -> str:
    root_m = str(root).replace("'", "''"); scripts_m = str(scripts_dir).replace("'", "''"); interface_m = str(interface).replace("'", "''"); model_m = model.replace("'", "''"); mat_m = mat_name.replace("'", "''")
    return f"rootDir='{root_m}'; model='{model_m}'; initScripts={matlab_cell(init_scripts)}; addpath('{scripts_m}'); setup_ut_support(rootDir,initScripts); load_system(fullfile(rootDir,'{model_m}.slx')); ins=find_system(model,'SearchDepth',1,'BlockType','Inport'); outs=find_system(model,'SearchDepth',1,'BlockType','Outport'); inputNames=reshape(cellstr(string(get_param(ins,'Name'))),1,[]); outputNames=reshape(cellstr(string(get_param(outs,'Name'))),1,[]); p=struct('schema','tcsd-model-interface/v1','inputs',{{inputNames}},'outputs',{{outputNames}}); fid=fopen('{interface_m}','w'); fprintf(fid,'%s',jsonencode(p,PrettyPrint=true)); fclose(fid); trace_logical_mcdc(rootDir,{{model}},'{mat_m}','WorkspaceInitialized',true); bdclose(model);"
def interface_names(values: Any) -> list[str]:
    if values is None: return []
    if isinstance(values, (str, int, float)): values = [values]
    if isinstance(values, dict): values = [values]
    return [str(item.get("name")) if isinstance(item, dict) else str(item) for item in values]
def validate_interface(value: dict[str, Any]) -> dict[str, Any]:
    if value.get("schema") != "tcsd-model-interface/v1": raise RuntimeError("model interface schema is invalid")
    if not all(isinstance(value.get(key), list) and all(isinstance(name, str) and name for name in value[key]) for key in ("inputs", "outputs")): raise RuntimeError("model interface inputs/outputs must be string arrays")
    return value
def initial_spec(interface: dict[str, Any], model: str) -> dict[str, Any]:
    root = interface.get("rootPorts", interface); inputs = root.get("inputs", []); outputs_ = root.get("outputs", [])
    input_names, output_names = interface_names(inputs), interface_names(outputs_)
    initialization = "\n".join(f"{name}=0;" for name in input_names)
    # TCSD Action 语法是 marker-first：输入赋值必须位于某个步标记之后，
    # 且末尾标记是纯观察窗口，不携带任何期望值。赋值写在首个标记之前会被
    # extract/backfill 的解析器静默丢弃，导致基线被回填成空延时步而缺失
    # expValue（Stage 08 硬失败）。因此基线用例将赋值作为第 1 步的刺激主体，
    # 并以独立观察窗口收尾。
    action = "\n".join(["[+0.1s]", *(f"{name}=0;" for name in input_names), "[+0.1s] // final observation window"])
    return {"model_name": model, "test_group": {"id": "TG_001", "name": model, "description": "确定性覆盖率基线"}, "tests": [{"id": "TC_001", "name": "确定性基线", "description": "由模型接口生成的确定性基线", "initialization": initialization, "action": action}]}
STEP_MARKER_RE = re.compile(r"^\s*\[\+")
EXP_VALUE_RE = re.compile(r"^\s*([A-Za-z_]\w*)\s*=\s*expValue\(\s*([-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)\s*\)\s*;?\s*$")
def workbook_steps(action: str) -> list[dict[str, Any]]:
    steps: list[dict[str, Any]] = []; current: dict[str, Any] | None = None; prelude: list[str] = []
    for raw in (action or "").splitlines():
        if STEP_MARKER_RE.match(raw):
            if current is not None: steps.append(current)
            current = {"lines": list(prelude)}
            prelude = []
        elif current is not None: current["lines"].append(raw)
        else: prelude.append(raw)
    if current is not None: steps.append(current)
    for index, step in enumerate(steps, 1):
        values: dict[str, float] = {}; non_expected = []
        for raw in step["lines"]:
            match = EXP_VALUE_RE.match(raw)
            if match:
                if match.group(1) in values: raise RuntimeError(f"duplicate workbook expValue at step {index}: {match.group(1)}")
                values[match.group(1)] = float(match.group(2))
            elif raw.strip(): non_expected.append(raw)
        step.update({"index": index, "values": values, "finalEmptyDelay": index == len(steps) and not non_expected})
    return steps
def simulation_backfill_evidence(simulation: dict[str, Any], workbook: Path) -> dict[str, Any]:
    from openpyxl import load_workbook
    tests = simulation.get("tests", []); tests = [tests] if isinstance(tests, dict) else tests
    simulated: dict[tuple[int, str], dict[int, dict[str, Any]]] = {}
    for case in tests:
        key = (int(case.get("row") or 0), str(case.get("test_id") or ""))
        if not key[0] or not key[1] or key in simulated: raise RuntimeError(f"invalid or duplicate simulation case identity: {key}")
        steps = case.get("steps", []); steps = [steps] if isinstance(steps, dict) else steps; indexed: dict[int, dict[str, Any]] = {}
        for step in steps:
            index = int(step.get("index") or 0)
            if not index or index in indexed or not isinstance(step.get("outputs", {}), dict): raise RuntimeError(f"invalid or duplicate simulation step: {key} step {index}")
            indexed[index] = step
        simulated[key] = indexed
    wb = load_workbook(workbook, read_only=True, data_only=False); ws = wb["TCSD"]; workbook_cases: dict[tuple[int, str], list[dict[str, Any]]] = {}
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 3).value == "Test": workbook_cases[(row, str(ws.cell(row, 1).value or ""))] = workbook_steps(str(ws.cell(row, 7).value or ""))
    wb.close()
    if set(simulated) != set(workbook_cases): raise RuntimeError(f"simulation/workbook case identity mismatch: simulation={sorted(simulated)}, workbook={sorted(workbook_cases)}")
    matched: list[dict[str, Any]] = []; case_outputs: dict[str, dict[str, int]] = {}
    for key, steps in workbook_cases.items():
        simulation_steps = simulated[key]
        if set(simulation_steps) != {step["index"] for step in steps}: raise RuntimeError(f"simulation/workbook step mismatch: {key}")
        counts: dict[str, int] = {}
        for step in steps:
            result = simulation_steps[step["index"]]; stable = result.get("stable", {}); outputs_ = result.get("outputs", {})
            expected = {} if step["finalEmptyDelay"] else {str(name): value for name, value in outputs_.items() if stable.get(name) is not False}
            actual = step["values"]
            if set(expected) != set(actual): raise RuntimeError(f"simulation/workbook output mismatch: {key} step {step['index']} expected={sorted(expected)} actual={sorted(actual)}")
            for name, value in expected.items():
                if isinstance(value, (dict, list)) or not math.isclose(float(value), actual[name], rel_tol=1e-7, abs_tol=1e-7): raise RuntimeError(f"simulation/workbook value mismatch: {key} step {step['index']} output {name}")
                counts[name] = counts.get(name, 0) + 1; matched.append({"row": key[0], "testId": key[1], "step": step["index"], "output": name, "value": actual[name]})
        case_outputs[f"{key[0]}:{key[1]}"] = counts
    if not matched: raise RuntimeError("simulation backfill produced no verified expValue items")
    return {"simulationValueCount": len(matched), "workbookBackfillCount": len(matched), "caseOutputCounts": case_outputs, "backfillItems": matched}
def coverage_meets(report: dict[str, Any], threshold: float) -> bool:
    records = report.get("models", report)
    valid = [record for record in records.values() if isinstance(record, dict) and all(key in record for key in ("condition", "decision", "mcdc"))]
    return bool(valid) and all(float(record[key]["percent"]) >= threshold for record in valid for key in ("condition", "decision", "mcdc"))

def stage_run(
    stage: int,
    job: dict[str, Any],
    *,
    stage10_mode: str = "auto",
    repair_brief: str = "",
    repair_proposal: str = "",
) -> None:
    root, out, model, state = workspace(job), outputs(job), model_name(job), load_state(job); inp = job["input"]; out.mkdir(parents=True, exist_ok=True)
    quality = load_module("tcsd_quality", scripts() / "run_tcsd_quality_loop.py")
    if stage == 1:
        required = [ensure_within(root, Path(inp["modelSlxPath"]), "modelSlxPath"), ensure_within(root, Path(inp["modelMatPath"]), "modelMatPath")]
        required.extend(ensure_within(root, root / item, "projectInitScript") for item in inp.get("projectInitScripts", []))
        ensure_within(root, out, "outputDir")
        missing = [str(item) for item in required if not item.is_file()]
        if missing: raise RuntimeError(f"required inputs missing: {missing}")
        manifest = out / ".tcsd-evidence" / "input-manifest.json"; write_json(manifest, {"schema": "tcsd-input-manifest/v1", "jobId": job["jobId"], "files": [{"path": str(item), "size": item.stat().st_size} for item in required], "projectAddon": inp.get("projectAddonCopy", {})})
        finish(job, stage, summary="输入文件与项目附件已验证。", artifacts=[artifact(root, manifest)]); return
    if stage == 2:
        # Gateway topology: MATLAB lives on the host behind the Gateway, so the
        # local executable check is skipped; the canary below verifies Gateway
        # health and the controlled transport instead.
        gateway_mode = bool(os.environ.get("TCSD_GATEWAY_TRANSPORT", "").strip()) and bool(os.environ.get("SATK_GATEWAY_URL", "").strip())
        if not gateway_mode:
            matlab_root = matlab_root_path(inp)
            matlab = matlab_root / "bin" / ("matlab.exe" if os.name == "nt" else "matlab")
            if not matlab.exists(): raise RuntimeError(f"MATLAB executable missing: {matlab}")
        if not (scripts() / "satk_eval.py").is_file(): raise RuntimeError("SATK runtime runner is missing")
        env = out / ".tcsd-evidence" / "environment.json"; env.parent.mkdir(parents=True, exist_ok=True)
        fixture = os.environ.get("TCSD_PIPELINE_ENV_CANARY_FIXTURE", "")
        if fixture:
            shutil.copy2(fixture, env)
        else:
            satk_runtime = load_module("tcsd_satk_eval", scripts() / "satk_eval.py")
            if os.environ.get("TCSD_GATEWAY_TRANSPORT", "").strip() and satk_runtime.gateway_url():
                selected_server = {"discovery": "matlab-gateway", "transport": "tcsd-gateway-transport"}
            else:
                selected_server = satk_runtime.server_info()
            modules: dict[str, dict[str, str]] = {}
            for module_name in ("yaml", "openpyxl"):
                module = __import__(module_name)
                version = str(getattr(module, "__version__", "") or "")
                if not version: raise RuntimeError(f"Python dependency has no version: {module_name}")
                modules[module_name] = {"version": version}
            io_dir = out / ".tcsd-runtime"; io_dir.mkdir(parents=True, exist_ok=True)
            io_sentinel = io_dir / f"environment-io-{secrets.token_hex(8)}.txt"
            io_value = secrets.token_hex(16)
            io_sentinel.write_text(io_value, encoding="utf-8")
            io_matched = io_sentinel.read_text(encoding="utf-8") == io_value
            io_sentinel.unlink()
            if not io_matched or io_sentinel.exists(): raise RuntimeError("workspace create/read/delete sentinel failed")
            nonce = secrets.token_hex(20)
            canary_script = io_dir / "stage02_environment_canary.m"
            matlab_sentinel = out / ".tcsd-evidence" / "matlab-satk-canary.json"
            canary_script.write_text(
                "\n".join([
                    f"nonce={matlab_string(nonce)};",
                    "assert(license('test','Simulink') == 1, 'Simulink license is unavailable');",
                    "matlabInfo=ver('MATLAB'); simulinkInfo=ver('Simulink');",
                    "assert(~isempty(matlabInfo) && ~isempty(simulinkInfo), 'MATLAB or Simulink version is unavailable');",
                    "load_system('simulink'); simulinkLoaded=bdIsLoaded('simulink'); bdclose('simulink');",
                    "assert(simulinkLoaded, 'Simulink library did not load');",
                    (
                        "p=struct('schema','tcsd-matlab-satk-canary/v1','nonce',nonce,"
                        "'matlabVersion',matlabInfo(1).Version,'simulinkVersion',simulinkInfo(1).Version,"
                        "'licenseAvailable',true,'simulinkLoaded',true);"
                    ),
                    f"fid=fopen({matlab_string(str(matlab_sentinel))},'w');",
                    "assert(fid >= 0, 'Unable to open the TCSD canary sentinel');",
                    "fprintf(fid,'%s',jsonencode(p,PrettyPrint=true)); fclose(fid);",
                ]),
                encoding="utf-8",
            )
            run_satk(
                [sys.executable, str(scripts() / "satk_eval.py"), str(canary_script)],
                root,
                stage=stage,
                context="environment gate failed",
            )
            if not matlab_sentinel.is_file(): raise RuntimeError("SATK/MCP returned without writing the MATLAB sentinel")
            matlab_result = read_json(matlab_sentinel)
            if (
                matlab_result.get("schema") != "tcsd-matlab-satk-canary/v1"
                or matlab_result.get("nonce") != nonce
                or matlab_result.get("licenseAvailable") is not True
                or matlab_result.get("simulinkLoaded") is not True
            ): raise RuntimeError("MATLAB/SATK canary sentinel is invalid or has the wrong nonce")
            write_json(env, {
                "schema": "tcsd-environment-gate/v2",
                "jobId": job["jobId"],
                "nonce": nonce,
                "matlabRoot": "host-matlab-gateway" if gateway_mode else str(matlab_root),
                "python": sys.executable,
                "runner": str(scripts() / "satk_eval.py"),
                "pythonDependencies": {"passed": True, "modules": modules},
                "workspaceIo": {"passed": True, "created": True, "readMatched": True, "deleted": True},
                "matlab": {"passed": True, "nonce": nonce, "version": str(matlab_result["matlabVersion"])},
                "simulink": {
                    "passed": True,
                    "licenseAvailable": True,
                    "loaded": True,
                    "version": str(matlab_result["simulinkVersion"]),
                },
                "satkMcp": {
                    "passed": True,
                    "runner": str(scripts() / "satk_eval.py"),
                    "server": selected_server,
                    "sentinelWritten": True,
                    "nonceMatched": True,
                },
                "passed": True,
            })
        gate = read_json(env)
        if gate.get("schema") != "tcsd-environment-gate/v2" or gate.get("jobId") != job["jobId"] or gate.get("passed") is not True:
            raise RuntimeError("environment canary fixture/result is invalid")
        finish(job, stage, summary="Python、工作目录、MATLAB、Simulink 与 SATK/MCP 执行门禁通过。", artifacts=[artifact(root, env)]); return
    if stage == 3:
        runtime = out / ".tcsd-runtime"; runtime.mkdir(exist_ok=True); resources = runtime / "owned-resources.json"
        init_manifest = out / ".tcsd-evidence" / "workspace-initialization.json"; entry = runtime / "stage03_initialize.m"; init = inp.get("projectInitScripts", [])
        root_m = str(root).replace("'", "''"); scripts_m = str(scripts()).replace("'", "''"); manifest_m = str(init_manifest).replace("'", "''")
        job_id_m = str(job["jobId"]).replace("'", "''")
        entry.write_text(f"rootDir='{root_m}'; initScripts={matlab_cell(init)}; addpath('{scripts_m}'); setup_ut_support(rootDir,initScripts); p=struct('schema','tcsd-workspace-initialization/v1','jobId','{job_id_m}','workspace',rootDir,'initScripts',{{initScripts}},'completed',true); fid=fopen('{manifest_m}','w'); fprintf(fid,'%s',jsonencode(p,PrettyPrint=true)); fclose(fid);", encoding="utf-8")
        if os.environ.get("TCSD_PIPELINE_SETUP_FIXTURE") == "1": write_json(init_manifest, {"schema":"tcsd-workspace-initialization/v1","jobId":job["jobId"],"workspace":str(root),"initScripts":init,"completed":True})
        else:
            run_satk(
                [sys.executable, str(scripts() / "satk_eval.py"), str(entry)],
                root,
                stage=stage,
                context="workspace initialization failed",
            )
        initialized = read_json(init_manifest)
        if initialized.get("jobId") != job["jobId"] or initialized.get("completed") is not True: raise RuntimeError("workspace initialization manifest is invalid")
        write_json(resources, {"schema": "tcsd-owned-resources/v1", "jobId": job["jobId"], "workspace": str(root), "generatedEntries": [str(entry.relative_to(root))]}); state["resources"] = str(resources); state["initializationManifest"] = str(init_manifest); save_state(job, state)
        finish(job, stage, summary="模型工作区已真实初始化并登记 job 资源所有权。", artifacts=[artifact(root, init_manifest), artifact(root, resources)], evidence={"initializationManifest":str(init_manifest.relative_to(root))}); return
    interface = out / f"{model}_interface.json"; traces = out / f"{model}_logical_traces.json"
    if stage == 4:
        entry = out / ".tcsd-runtime" / "stage04_interface.m"
        code = stage4_matlab_code(root=root, scripts_dir=scripts(), interface=interface, model=model, mat_name=Path(inp["modelMatPath"]).name, init_scripts=inp.get("projectInitScripts", []))
        entry.write_text(code, encoding="utf-8")
        run_satk(
            [sys.executable, str(scripts() / "satk_eval.py"), str(entry)],
            root,
            stage=stage,
            context="model interface extraction failed",
        )
        validate_interface(read_json(interface)); read_json(traces); state.update({"interface": str(interface), "traces": str(traces)}); save_state(job, state)
        finish(job, stage, summary="模型已加载并提取根输入输出接口。", artifacts=[artifact(root, interface), artifact(root, traces)]); return
    mapping, obligations, coverage_ir = out / f"{model}_logical_operators.json", out / f"{model}_coverage_obligations.json", out / f"{model}_coverage_ir.json"
    if stage == 5:
        run([sys.executable, str(scripts()/"derive_logical_mcdc_mappings.py"), "--traces", str(traces), "--output", str(mapping)], root)
        run([sys.executable, str(scripts()/"build_logical_mcdc_obligations.py"), "--logical-operators", str(mapping), "--output", str(obligations), "--allow-unresolved"], root)
        decision_blocks = out / f"{model}_decision_blocks.json"
        decision_obligations = out / f"{model}_decision_obligations.json"
        root_m = str(root).replace("'", "''"); scripts_m = str(scripts()).replace("'", "''")
        model_m = model.replace("'", "''"); mat_m = Path(inp["modelMatPath"]).name.replace("'", "''")
        init = inp.get("projectInitScripts", [])
        collected = False
        # Prefer MATLAB-collected decision blocks (real connectivity evidence);
        # fall back to static SLX XML analysis when the MATLAB gate fails.
        try:
            entry = out / ".tcsd-runtime" / "stage05_decision_blocks.m"
            entry.write_text(
                f"rootDir='{root_m}'; initScripts={matlab_cell(init)}; addpath('{scripts_m}'); "
                f"setup_ut_support(rootDir,initScripts); "
                f"collect_decision_blocks(rootDir,'{model_m}','{mat_m}',initScripts,'{str(decision_blocks).replace(chr(39), chr(39)+chr(39))}');",
                encoding="utf-8",
            )
            run_satk([sys.executable, str(scripts() / "satk_eval.py"), str(entry)], root, stage=stage, context="decision blocks collection failed")
            if decision_blocks.is_file():
                run([sys.executable, str(scripts()/"build_decision_obligations.py"), "--blocks", str(decision_blocks), "--slx", str(inp["modelSlxPath"]), "--output", str(decision_obligations)], root)
                collected = True
        except BaseException as error:
            print(f"stage 05 MATLAB decision blocks collection unavailable: {error}", file=sys.stderr)
        if not collected:
            try:
                run([sys.executable, str(scripts()/"build_decision_obligations.py"), "--slx", str(inp["modelSlxPath"]), "--interface", str(interface), "--output", str(decision_obligations)], root)
                collected = True
            except BaseException as error:
                print(f"stage 05 static decision obligations unavailable: {error}", file=sys.stderr)
        extra_args = ["--decision-obligations", str(decision_obligations)] if collected and decision_obligations.is_file() else []
        run([sys.executable, str(scripts()/"build_coverage_ir.py"), "--logical-traces", str(traces), "--obligations", str(obligations), *extra_args, "--output", str(coverage_ir)], root)
        state.update({"mapping": str(mapping), "obligations": str(obligations), "coverageIr": str(coverage_ir), "decisionObligations": str(decision_obligations)}); save_state(job, state)
        finish(job, stage, summary="Condition、Decision 与 MC/DC 覆盖目标已形成 Coverage IR。", artifacts=[artifact(root, mapping), artifact(root, obligations), artifact(root, coverage_ir), artifact(root, decision_obligations)]); return
    if stage == 6:
        plan = out / f"{model}_state_probe_plan.json"
        capability_path, precheck_evidence = run_probe_capability_precheck(job, model, root, out, mapping)
        plan_cmd = [sys.executable, str(scripts()/"build_state_probe_plan.py"), "--traces", str(traces), "--output", str(plan)]
        if capability_path is not None:
            plan_cmd += ["--probe-capability", str(capability_path)]
        run(plan_cmd, root); plan_data = read_json(plan)
        probe_artifacts = [artifact(root, plan)]
        if capability_path is not None:
            probe_artifacts.append(artifact(root, capability_path))
        candidate_count = int(plan_data.get("summary", {}).get("candidate_count") or len(plan_data.get("tests", [])))
        unprobeable_target_count = int(plan_data.get("summary", {}).get("unprobeable_target_count") or 0)
        probe_timeout_seconds = stage6_probe_timeout_seconds(candidate_count)
        if candidate_count > 0:
            probe_results = out / f"{model}_state_probe_results.json"; probe_fixture = os.environ.get("TCSD_PIPELINE_PROBE_RESULTS_FIXTURE", "")
            if probe_fixture:
                shutil.copy2(probe_fixture, probe_results); run([sys.executable, str(scripts()/"build_probe_mcdc_obligations.py"), "--probe-results", str(probe_results), "--model", model, "--output-dir", str(out), "--logical-mappings", str(mapping)], root)
            else: obligations, _ = quality.run_probe(python=sys.executable, scripts=scripts(), root_dir=root, model=model, mat_file=inp["modelMatPath"], init_scripts=inp.get("projectInitScripts", []), unreachable_overrides="", collect_coverage=False, coverage_threshold=float(inp.get("coverageThreshold", 80)), case_json=plan, output_name=f"{model}_state_probe_results.json", gateway_timeout_seconds=probe_timeout_seconds)
            read_json(probe_results)
            ir_args = [sys.executable, str(scripts()/"build_coverage_ir.py"), "--logical-traces", str(traces), "--probe-results", str(probe_results), "--obligations", str(obligations)]
            if state.get("decisionObligations") and Path(state["decisionObligations"]).is_file():
                ir_args += ["--decision-obligations", str(state["decisionObligations"])]
            run(ir_args + ["--output", str(coverage_ir)], root); probe_artifacts.extend([artifact(root, probe_results), artifact(root, obligations), artifact(root, coverage_ir)])
        reconciliation = probe_observation_reconciliation(probe_results) if candidate_count > 0 else None
        gap_count = int(reconciliation["gapCount"]) if reconciliation else 0
        # Unprobeable targets are registered gaps too: with every target
        # unobservable the stage must NOT report completed (review P1).
        gap_count += unprobeable_target_count
        state["statePlan"] = str(plan)
        if reconciliation is not None:
            state["stateProbeReconciliation"] = reconciliation
        save_state(job, state)
        if candidate_count == 0:
            summary = "未发现需要额外 Probe 的状态及时序候选。"
        elif gap_count == 0:
            summary = "状态及时序刺激已生成并由实际 Probe 验证。"
        else:
            summary = f"状态及时序刺激已生成；{gap_count} 项缺口（含 {unprobeable_target_count} 个不可探测目标）未获可信观测，已登记（partial），以第 9 阶段实测覆盖为准。"
        finish(
            job, stage,
            status="partial" if gap_count > 0 else "completed",
            summary=summary,
            artifacts=probe_artifacts,
            evidence={
                "candidateCount": candidate_count,
                "probeExecuted": candidate_count > 0,
                "probeTimeoutSeconds": probe_timeout_seconds if candidate_count > 0 else None,
                "unprobeableTargetCount": unprobeable_target_count,
                "precheck": precheck_evidence,
                "reconciliation": reconciliation,
            },
        ); return
    spec, workbook = out / f"{model}_tcsd_spec.json", out / f"{model}_Test0001_tcsd.xlsx"
    if stage == 7:
        write_json(spec, initial_spec(read_json(interface), model)); run([sys.executable, str(scripts()/"build_tcsd_from_json.py"), "--template", str(scripts().parent/"assets"/"templates"/"tcsd_template.xlsx"), "--spec", str(spec), "--output", str(workbook), "--interface-json", str(interface)], root)
        spec, workbook, _ = quality.synthesize_ir_once(python=sys.executable, scripts=scripts(), root_dir=root, template=scripts().parent/"assets"/"templates"/"tcsd_template.xlsx", model=model, spec=spec, workbook=workbook, interface_json=interface, coverage_ir=coverage_ir, iteration=0)
        quality.validate_workbook(python=sys.executable, scripts=scripts(), root_dir=root, workbook=workbook, interface_json=interface); state.update({"spec": str(spec), "workbook": str(workbook)}); save_state(job, state)
        planning_obligations = out / f"{model}_planning_obligations_snapshot.json"
        planning_assessment = out / f"{model}_planning_mapping_assessment.json"
        shutil.copy2(obligations, planning_obligations)
        assessment = quality.validate_mapping(
            python=sys.executable,
            scripts=scripts(),
            root_dir=root,
            workbook=workbook,
            obligations=planning_obligations,
            report=planning_assessment,
        )
        assessment = planning_mapping_assessment(assessment, planning_obligations, root)
        write_json(planning_assessment, assessment)
        finish(
            job,
            stage,
            summary="首版 TCSD 已生成并通过接口与工作簿校验；静态映射仅作为规划诊断。",
            artifacts=[
                artifact(root, spec),
                artifact(root, workbook, "xlsx", "workbook"),
                artifact(root, planning_obligations, "json", "planning-obligations"),
                artifact(root, planning_assessment, "json", "planning-mapping-assessment"),
            ],
            evidence={
                "planningMappingAssessment": str(planning_assessment.relative_to(root)),
                "mappingAuthority": "planning",
                "supersededByStage": 9,
            },
        ); return
    workbook = Path(state["workbook"]); spec = Path(state["spec"]); threshold = float(inp.get("coverageThreshold", 80))
    if stage == 8:
        cases = quality.extract_cases(python=sys.executable, scripts=scripts(), root_dir=root, model=model, workbook=workbook, interface_json=interface)
        sim = quality.simulate_and_backfill(python=sys.executable, scripts=scripts(), root_dir=root, model=model, workbook=workbook, case_json=cases, mat_file=inp["modelMatPath"], outputs=",".join(read_json(interface).get("outputs", [])), exclude_outputs="", interface_json=interface)
        backfill = simulation_backfill_evidence(read_json(sim), workbook)
        state.update({"cases": str(cases), "initialSimulation": str(sim), "initialBackfillEvidence": backfill, "expValueCount": backfill["workbookBackfillCount"]}); save_state(job, state)
        finish(job, stage, summary="首版仿真完成，expValue 已由实际仿真回填。", artifacts=[artifact(root, workbook, "xlsx", "workbook"), artifact(root, sim)], evidence={"simulationResult": str(sim.relative_to(root)), "expValueCount": backfill["workbookBackfillCount"], **backfill}); return
    initial_cov = out / f"{model}_initial_coverage_summary.json"
    if stage == 9:
        ob, cov = quality.run_probe(python=sys.executable, scripts=scripts(), root_dir=root, model=model, mat_file=inp["modelMatPath"], init_scripts=inp.get("projectInitScripts", []), unreachable_overrides="", collect_coverage=True, coverage_threshold=threshold); report={"schema":"tcsd-coverage-report/v1","models":read_json(cov)}; write_json(initial_cov,report)
        # Rebuild the IR with the probe-refreshed obligations so the measured
        # observation vectors flow back into the stage-10 repair brief.
        if state.get("decisionObligations") and Path(state["decisionObligations"]).is_file():
            try:
                probe_results = out / "logic_probe_results.json"
                ir_args = [sys.executable, str(scripts()/"build_coverage_ir.py"), "--logical-traces", str(traces), "--probe-results", str(probe_results), "--obligations", str(ob), "--decision-obligations", str(state["decisionObligations"])]
                run(ir_args + ["--output", str(coverage_ir)], root)
                state["coverageIr"] = str(coverage_ir)
            except BaseException as error:
                print(f"stage 09 IR refresh unavailable: {error}", file=sys.stderr)
        state.update({"obligations":str(ob),"initialCoverage":str(initial_cov),"coverage":str(cov)}); save_state(job,state)
        finish(job,stage,summary="首轮 Condition、Decision 与 MC/DC 覆盖率已采集。",artifacts=[artifact(root,initial_cov)],coverage=report); return
    if stage == 10:
        report = read_json(initial_cov)
        brief = Path(repair_brief).resolve() if repair_brief else out / f"{model}_coverage_repair_brief.json"
        proposal = Path(repair_proposal).resolve() if repair_proposal else out / f"{model}_agent_coverage_repair_proposal.json"
        if stage10_mode in {"prepare", "auto"}:
            run(
                [
                    sys.executable,
                    str(scripts() / "validate_agent_coverage_repair.py"),
                    "prepare",
                    "--job-id",
                    str(job["jobId"]),
                    "--model",
                    model,
                    "--coverage-report",
                    str(initial_cov),
                    "--logical-traces",
                    str(traces),
                    "--coverage-ir",
                    str(coverage_ir),
                    "--interface",
                    str(interface),
                    "--threshold",
                    str(threshold),
                    "--probe-results",
                    str(out / f"{model}_state_probe_results.json"),
                    "--output",
                    str(brief),
                ],
                root,
            )
            if stage10_mode == "prepare":
                return
        if coverage_meets(report, threshold):
            finish(
                job,
                stage,
                status="skipped",
                summary=f"首轮三项覆盖率均达到 {threshold:g}%。",
                skipReason=f"首轮三项覆盖率均达到 {threshold:g}%。",
                artifacts=[artifact(root, brief)],
                evidence={"repairBrief": str(brief.relative_to(root))},
            )
            return
        if not proposal.is_file():
            raise RuntimeError("stage 10 requires an Agent-authored coverage repair proposal")

        proposal_ir = out / f"{model}_agent_repair_coverage_ir.json"
        proposal_validation = out / f"{model}_agent_repair_validation.json"
        run(
            [
                sys.executable,
                str(scripts() / "validate_agent_coverage_repair.py"),
                "validate",
                "--brief",
                str(brief),
                "--proposal",
                str(proposal),
                "--interface",
                str(interface),
                "--output-ir",
                str(proposal_ir),
                "--report-json",
                str(proposal_validation),
            ],
            root,
        )
        validation = read_json(proposal_validation)
        base_artifacts = [
            artifact(root, brief, "json", "coverage-repair-brief"),
            artifact(root, proposal, "json", "agent-repair-proposal"),
            artifact(root, proposal_validation, "json", "agent-repair-validation"),
            artifact(root, proposal_ir, "json", "coverage-repair"),
        ]
        accepted = int(validation.get("acceptedCandidateCount") or 0)
        unresolved = int(validation.get("unresolvedCount") or 0)
        evidence = {
            "repairBrief": str(brief.relative_to(root)),
            "repairProposal": str(proposal.relative_to(root)),
            "proposalValidation": str(proposal_validation.relative_to(root)),
            "coverageIr": str(proposal_ir.relative_to(root)),
            "proposalItemCount": int(validation.get("proposalItemCount") or 0),
            "acceptedCandidateCount": accepted,
            "unresolvedCount": unresolved,
        }
        if accepted == 0:
            reason = "agent_reported_specific_unresolved_deficits"
            state.update(
                {
                    "repairAttempted": True,
                    "repairApplied": False,
                    "repairReason": reason,
                    "repairEvidence": str(proposal_validation),
                }
            )
            save_state(job, state)
            finish(
                job,
                stage,
                status="partial",
                summary="Agent 已完成局部模型分析，但所有缺口均有具体的未解析原因，未追加无效用例。",
                artifacts=base_artifacts,
                repair={
                    "required": True,
                    "attempted": True,
                    "applied": False,
                    "passes": 0,
                    "reason": reason,
                    "evidence": str(proposal_validation.relative_to(root)),
                },
                evidence=evidence,
            )
            return

        next_spec, next_book, synthesis = quality.synthesize_ir_once(
            python=sys.executable,
            scripts=scripts(),
            root_dir=root,
            template=scripts().parent / "assets" / "templates" / "tcsd_template.xlsx",
            model=model,
            spec=spec,
            workbook=workbook,
            interface_json=interface,
            coverage_ir=proposal_ir,
            iteration=1,
        )
        synthesis_evidence = out / f"{model}_coverage_ir_synthesis_iter1.json"
        added = int(synthesis.get("added") or 0)
        evidence.update(
            {
                "synthesisReport": str(synthesis_evidence.relative_to(root)),
                "synthesisAddedCount": added,
            }
        )
        repair_artifacts = [*base_artifacts, artifact(root, synthesis_evidence)]
        if added <= 0:
            reason = "agent_candidates_duplicate_existing_tests"
            state.update(
                {
                    "repairAttempted": True,
                    "repairApplied": False,
                    "repairReason": reason,
                    "repairEvidence": str(synthesis_evidence),
                }
            )
            save_state(job, state)
            finish(
                job,
                stage,
                status="partial",
                summary="Agent 设计的候选与已有用例重复，宿主未追加重复测试。",
                artifacts=repair_artifacts,
                repair={
                    "required": True,
                    "attempted": True,
                    "applied": False,
                    "passes": 0,
                    "reason": reason,
                    "evidence": str(synthesis_evidence.relative_to(root)),
                },
                evidence=evidence,
            )
            return

        quality.validate_workbook(
            python=sys.executable,
            scripts=scripts(),
            root_dir=root,
            workbook=next_book,
            interface_json=interface,
        )
        candidate_diagnostic = out / f"{model}_repair_candidate_validation.json"
        try:
            candidate_cases = quality.extract_cases(
                python=sys.executable,
                scripts=scripts(),
                root_dir=root,
                model=model,
                workbook=next_book,
                interface_json=interface,
            )
            candidate_simulation = quality.simulate_and_backfill(
                python=sys.executable,
                scripts=scripts(),
                root_dir=root,
                model=model,
                workbook=next_book,
                case_json=candidate_cases,
                mat_file=inp["modelMatPath"],
                outputs=",".join(read_json(interface).get("outputs", [])),
                exclude_outputs="",
                interface_json=interface,
                result_name=f"{model}_repair_candidate_simulation.json",
            )
            backfill = simulation_backfill_evidence(read_json(candidate_simulation), next_book)
            write_json(
                candidate_diagnostic,
                {
                    "schema": "tcsd-repair-candidate-validation/v1",
                    "jobId": job["jobId"],
                    "passed": True,
                    "candidateCount": added,
                    "simulationResult": str(candidate_simulation.relative_to(root)),
                    "backfill": backfill,
                },
            )
        except Exception as error:
            reason = "agent_candidate_simulation_failed"
            write_json(
                candidate_diagnostic,
                {
                    "schema": "tcsd-repair-candidate-validation/v1",
                    "jobId": job["jobId"],
                    "passed": False,
                    "candidateCount": added,
                    "errorType": type(error).__name__,
                    "errorMessage": str(error),
                },
            )
            state.update(
                {
                    "repairAttempted": True,
                    "repairApplied": False,
                    "repairReason": reason,
                    "repairEvidence": str(candidate_diagnostic),
                }
            )
            save_state(job, state)
            finish(
                job,
                stage,
                status="partial",
                summary="Agent 候选未通过宿主仿真验证，已保留首轮工作簿。",
                artifacts=[*repair_artifacts, artifact(root, candidate_diagnostic)],
                repair={
                    "required": True,
                    "attempted": True,
                    "applied": False,
                    "passes": 0,
                    "reason": reason,
                    "evidence": str(candidate_diagnostic.relative_to(root)),
                },
                evidence={**evidence, "candidateValidation": str(candidate_diagnostic.relative_to(root))},
            )
            return

        evidence.update(
            {
                "candidateValidation": str(candidate_diagnostic.relative_to(root)),
                "candidateSimulation": str(candidate_simulation.relative_to(root)),
                "simulationResult": str(candidate_simulation.relative_to(root)),
                "candidateValidationPassed": True,
                "expValueCount": backfill["workbookBackfillCount"],
                **backfill,
            }
        )
        state.update(
            {
                "repairAttempted": True,
                "repairApplied": True,
                "repairReason": "agent_targeted_candidates_validated_and_appended",
                "repairEvidence": str(candidate_diagnostic),
                "workbook": str(next_book),
                "spec": str(next_spec),
                "repairCandidateSimulation": str(candidate_simulation),
            }
        )
        save_state(job, state)
        finish(
            job,
            stage,
            summary="Agent 已针对覆盖缺口完成局部模型分析，候选用例通过宿主校验与仿真并已追加。",
            artifacts=[
                *repair_artifacts,
                artifact(root, candidate_diagnostic),
                artifact(root, next_book, "xlsx", "workbook"),
                artifact(root, candidate_simulation),
            ],
            repair={
                "required": True,
                "attempted": True,
                "applied": True,
                "passes": 1,
                "reason": "agent_targeted_candidates_validated_and_appended",
                "evidence": str(candidate_diagnostic.relative_to(root)),
            },
            evidence=evidence,
        )
        return
    final_cov=out/f"{model}_final_coverage_summary.json"
    if stage == 11:
        if not state.get("repairApplied"): finish(job,stage,status="skipped",summary="修正未实际应用，引用首轮仿真与覆盖率。",skipReason="修正未实际应用，引用首轮结果。",artifacts=[]); return
        workbook=Path(state["workbook"]); cases=quality.extract_cases(python=sys.executable,scripts=scripts(),root_dir=root,model=model,workbook=workbook,interface_json=interface); case_count=len(read_json(cases).get("tests",[])); probe_timeout_seconds=stage11_probe_timeout_seconds(case_count); sim=quality.simulate_and_backfill(python=sys.executable,scripts=scripts(),root_dir=root,model=model,workbook=workbook,case_json=cases,mat_file=inp["modelMatPath"],outputs=",".join(read_json(interface).get("outputs",[])),exclude_outputs="",interface_json=interface,result_name=f"{model}_final_simulation_results.json"); backfill=simulation_backfill_evidence(read_json(sim),workbook); ob,cov=quality.run_probe(python=sys.executable,scripts=scripts(),root_dir=root,model=model,mat_file=inp["modelMatPath"],init_scripts=inp.get("projectInitScripts",[]),unreachable_overrides="",collect_coverage=True,coverage_threshold=threshold,gateway_timeout_seconds=probe_timeout_seconds); final_report={"schema":"tcsd-coverage-report/v1","models":read_json(cov)}; write_json(final_cov,final_report); state.update({"finalSimulation":str(sim),"finalCoverage":str(final_cov),"finalBackfillEvidence":backfill,"obligations":str(ob)}); save_state(job,state)
        finish(job,stage,summary="修正后最终仿真、回填与覆盖率检查已完成。",artifacts=[artifact(root,workbook,"xlsx","workbook"),artifact(root,sim),artifact(root,final_cov)],coverage=final_report,evidence={"simulationResult":str(sim.relative_to(root)),"caseCount":case_count,"probeTimeoutSeconds":probe_timeout_seconds,"expValueCount":backfill["workbookBackfillCount"],**backfill}); return
    if stage == 12:
        cleanup=out/f"{model}_tcsd_cleanup.json"
        owned_candidates=[out/".tcsd-runtime"/"stage04_interface.m",out/".tcsd-runtime"/"job.json",out/f"{model}_probe_mcdc_entry.m",out/f"{model}_simulate_mcdc_entry.m"]; removed=[]
        for candidate in owned_candidates:
            resolved=candidate.resolve()
            if resolved.is_relative_to(root) and resolved.exists(): resolved.unlink(); removed.append(resolved.relative_to(root).as_posix())
        write_json(cleanup,{"schema":"tcsd-cleanup-result/v1","jobId":job["jobId"],"ownerJobId":job["jobId"],"closedMatlabSessions":[],"stoppedMcpProcesses":[],"removedEntries":removed})
        finish(job,stage,status="completed",summary="任务归属资源清理证据已生成，等待宿主整理最终 manifest。",artifacts=[artifact(root,cleanup)],evidence={"cleanup":str(cleanup.relative_to(root))}); return
    raise RuntimeError(f"unsupported stage {stage}")

def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--result", required=True)
    parser.add_argument("--stage10-mode", choices=("auto", "prepare", "apply"), default="auto")
    parser.add_argument("--repair-brief", default="")
    parser.add_argument("--repair-proposal", default="")
    args = parser.parse_args()
    manifest = read_json(Path(args.manifest))
    if manifest.get("schema") != INPUT_SCHEMA: raise RuntimeError("stage input manifest schema is invalid")
    stage = int(manifest.get("stageIndex") or 0)
    if stage < 1 or stage > 12: raise RuntimeError("stage input manifest stageIndex is invalid")
    job = manifest.get("job")
    if not isinstance(job, dict) or job.get("jobId") != manifest.get("jobId"): raise RuntimeError("stage input manifest job snapshot is invalid")
    result_path = ensure_within(workspace(job), Path(args.result), "resultPath")
    job["_stageResultPath"] = str(result_path)
    if args.repair_brief:
        ensure_within(workspace(job), Path(args.repair_brief), "repairBriefPath")
    if args.repair_proposal:
        ensure_within(workspace(job), Path(args.repair_proposal), "repairProposalPath")
    try:
        stage_run(
            stage,
            job,
            stage10_mode=args.stage10_mode,
            repair_brief=args.repair_brief,
            repair_proposal=args.repair_proposal,
        )
        if stage == 10 and args.stage10_mode == "prepare":
            if not args.repair_brief or not Path(args.repair_brief).is_file():
                raise RuntimeError("stage 10 repair brief was not created")
            return 0
        result = read_json(result_path)
        if result.get("schema") != RESULT_SCHEMA or result.get("jobId") != job["jobId"] or result.get("stageIndex") != stage:
            raise RuntimeError("stage runtime produced an invalid result envelope")
        return 0
    except BaseException as error:
        write_json(result_path, {
            "schema": RESULT_SCHEMA,
            "jobId": job["jobId"],
            "stageIndex": stage,
            "status": "failed",
            "summary": "TCSD deterministic stage runtime failed.",
            "artifacts": [],
            "error": {
                "code": hard_error_code(stage, error),
                "message": str(error),
                "hard": True
            }
        })
        return 1
if __name__ == "__main__": raise SystemExit(main())
