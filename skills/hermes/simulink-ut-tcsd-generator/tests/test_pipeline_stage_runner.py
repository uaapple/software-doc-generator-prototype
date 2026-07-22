import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "run_tcsd_pipeline_stage.py"
SPEC = importlib.util.spec_from_file_location("run_tcsd_pipeline_stage", SCRIPT)
RUNNER = importlib.util.module_from_spec(SPEC)
assert SPEC.loader
SPEC.loader.exec_module(RUNNER)


class PipelineStageRunnerTests(unittest.TestCase):
    def test_first_three_stages_write_owned_authoritative_checkpoints(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            output = root / "outputs"
            model = root / "GenericModel.slx"
            mat = root / "GenericModel.mat"
            model.write_bytes(b"slx")
            mat.write_bytes(b"mat")
            job = {
                "jobId": "job-generic",
                "events": [],
                "resources": {"ownerJobId": "job-generic"},
                "input": {
                    "workspaceDir": str(root),
                    "outputDir": str(output),
                    "modelSlxPath": str(model),
                    "modelMatPath": str(mat),
                    "projectAddonCopy": {"copiedFileCount": 0},
                },
            }
            old = os.environ.get("TCSD_PIPELINE_SKIP_MATLAB_GATE")
            old_setup = os.environ.get("TCSD_PIPELINE_SETUP_FIXTURE")
            os.environ["TCSD_PIPELINE_SKIP_MATLAB_GATE"] = "1"
            os.environ["TCSD_PIPELINE_SETUP_FIXTURE"] = "1"
            try:
                for stage in (1, 2, 3):
                    RUNNER.stage_run(stage, job)
                    checkpoint = json.loads((output / ".tcsd-checkpoints" / f"stage-{stage:02d}.json").read_text(encoding="utf-8"))
                    self.assertEqual(checkpoint["schema"], "tcsd-stage-checkpoint/v1")
                    self.assertEqual(checkpoint["jobId"], "job-generic")
                    self.assertEqual(checkpoint["stageIndex"], stage)
                    self.assertTrue(checkpoint["artifacts"])
                resources = json.loads((output / ".tcsd-runtime" / "owned-resources.json").read_text(encoding="utf-8"))
                self.assertEqual(resources["jobId"], "job-generic")
                initialized = json.loads((output / ".tcsd-checkpoints" / "workspace-initialization.json").read_text(encoding="utf-8"))
                self.assertEqual(initialized["schema"], "tcsd-workspace-initialization/v1")
                self.assertTrue(initialized["completed"])
            finally:
                if old is None:
                    os.environ.pop("TCSD_PIPELINE_SKIP_MATLAB_GATE", None)
                else:
                    os.environ["TCSD_PIPELINE_SKIP_MATLAB_GATE"] = old
                if old_setup is None:
                    os.environ.pop("TCSD_PIPELINE_SETUP_FIXTURE", None)
                else:
                    os.environ["TCSD_PIPELINE_SETUP_FIXTURE"] = old_setup

    def test_single_input_output_interface_is_not_split_into_characters(self):
        interface = RUNNER.validate_interface({"schema": "tcsd-model-interface/v1", "inputs": ["OnlyInput"], "outputs": ["OnlyOutput"]})
        spec = RUNNER.initial_spec({"inputs": "OnlyInput", "outputs": "OnlyOutput"}, "GenericModel")
        self.assertEqual(interface["inputs"], ["OnlyInput"])
        self.assertEqual(spec["tests"][0]["initialization"], "OnlyInput=0;")
        self.assertNotIn("output_reference", spec["tests"][0])
        with self.assertRaises(RuntimeError):
            RUNNER.validate_interface({"schema": "tcsd-model-interface/v1", "inputs": "OnlyInput", "outputs": ["OnlyOutput"]})

    def test_stage_four_initializes_its_new_session_before_tracing(self):
        code = RUNNER.stage4_matlab_code(root=Path("C:/job"), scripts_dir=Path("C:/skill/scripts"), interface=Path("C:/job/outputs/interface.json"), model="GenericModel", mat_name="GenericModel.mat", init_scripts=["project_init.m"])
        setup = code.index("setup_ut_support(rootDir,initScripts)")
        trace = code.index("trace_logical_mcdc")
        self.assertIn("initScripts={'project_init.m'}", code)
        self.assertIn("'WorkspaceInitialized',true", code)
        self.assertLess(setup, trace)

    def test_simulation_backfill_requires_matching_real_result_counts(self):
        with tempfile.TemporaryDirectory() as temp:
            workbook = Path(temp) / "result.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "TCSD"; ws["A3"] = "TC_001"; ws["C3"] = "Test"; ws["G3"] = "[+0.1s]\nOnlyOutput = expValue(1);\n[+0.1s]"; wb.save(workbook)
            simulation = {"tests": [{"row": 3, "test_id": "TC_001", "steps": [{"index": 1, "outputs": {"OnlyOutput": 1}, "stable": {"OnlyOutput": True}}, {"index": 2, "outputs": {"OnlyOutput": 1}, "stable": {"OnlyOutput": True}}]}]}
            evidence = RUNNER.simulation_backfill_evidence(simulation, workbook)
            self.assertEqual(evidence["simulationValueCount"], 1)
            self.assertEqual(evidence["workbookBackfillCount"], 1)
            self.assertEqual(evidence["backfillItems"], [{"row": 3, "testId": "TC_001", "step": 1, "output": "OnlyOutput", "value": 1.0}])
            with self.assertRaises(RuntimeError):
                RUNNER.simulation_backfill_evidence({"tests": []}, workbook)

    def test_simulation_backfill_rejects_missing_extra_and_wrong_values(self):
        with tempfile.TemporaryDirectory() as temp:
            workbook = Path(temp) / "result.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "TCSD"; ws["A4"] = "TC_002"; ws["C4"] = "Test"
            simulation = {"tests": [{"row": 4, "test_id": "TC_002", "steps": [{"index": 1, "outputs": {"Result": 2.5}, "stable": {"Result": True}}]}]}
            for action in ("[+0.1s]", "[+0.1s]\nResult = expValue(2.5);\nExtra = expValue(1);", "[+0.1s]\nResult = expValue(2.4);"):
                ws["G4"] = action; wb.save(workbook)
                with self.assertRaises(RuntimeError): RUNNER.simulation_backfill_evidence(simulation, workbook)

    def test_stage_five_to_six_cli_persists_unresolved_then_runs_probe(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); output = root / "outputs"; output.mkdir()
            model = root / "GenericModel.slx"; mat = root / "GenericModel.mat"; model.write_bytes(b"slx"); mat.write_bytes(b"mat")
            trace = {
                "schema": "simulink-ut-logical-mcdc-trace/v2", "model": "GenericModel", "operator_count": 1,
                "operators": [{"id": "GenericModel:1", "sid": "GenericModel:1", "block_path": "GenericModel/Decision", "operator": "AND", "ports": [
                    {"index": 1, "trace": {"kind": "relational", "operator": ">", "inputs": [{"trace": {"kind": "stateful", "block_type": "UnitDelay", "inputs": [{"trace": {"kind": "root_inport", "signal": "Enable"}}]}}, {"trace": {"kind": "constant", "value": "0"}}]}},
                    {"index": 2, "trace": {"kind": "root_inport", "signal": "Request"}},
                ]}],
            }
            (output / "GenericModel_logical_traces.json").write_text(json.dumps(trace), encoding="utf-8")
            probe = {"GenericModel": {"schema": "simulink-ut-logical-mcdc-probe/v2", "model": "GenericModel", "probes": [{"id": "GenericModel:1", "sid": "GenericModel:1", "block_path": "GenericModel/Decision", "operator": "AND", "port_names": ["u1", "u2"]}], "observations": []}}
            for index, label in enumerate(("TT", "FT", "TF"), 1):
                probe["GenericModel"]["observations"].append({"test_id": f"STATE_PROBE_{index:04d}", "row": index, "step_index": 2, "time_s": 0.1, "inputs": {"Enable": int(label[0] == "T"), "Request": int(label[1] == "T")}, "params": {}, "vectors": {"decision": {"id": "GenericModel:1", "label": label, "ok": True}}, "stimulus": {"initial_inputs": {"Enable": 0, "Request": 1}, "initial_params": {}, "steps": [{"index": 1, "delay_s": 0.01, "input_updates": {"Enable": 1}, "param_updates": {}}, {"index": 2, "delay_s": 0.1, "input_updates": {}, "param_updates": {}}], "evidence_step": 2}, "prediction_status": "observed"})
            probe_fixture = root / "probe-results.json"; probe_fixture.write_text(json.dumps(probe), encoding="utf-8")
            job = {"jobId": "job-cli", "resources": {"ownerJobId": "job-cli"}, "input": {"workspaceDir": str(root), "outputDir": str(output), "modelSlxPath": str(model), "modelMatPath": str(mat), "coverageThreshold": 80}}
            job_path = root / "job.json"; job_path.write_text(json.dumps(job), encoding="utf-8")
            subprocess.run([sys.executable, str(SCRIPT), "--job", str(job_path), "--stage", "5"], check=True, cwd=root)
            before = json.loads((output / "GenericModel_coverage_obligations.json").read_text(encoding="utf-8"))
            self.assertGreater(before["summary"]["unresolved_count"], 0)
            env = dict(os.environ); env["TCSD_PIPELINE_PROBE_RESULTS_FIXTURE"] = str(probe_fixture)
            subprocess.run([sys.executable, str(SCRIPT), "--job", str(job_path), "--stage", "6"], check=True, cwd=root, env=env)
            checkpoint = json.loads((output / ".tcsd-checkpoints" / "stage-06.json").read_text(encoding="utf-8"))
            after = json.loads((output / "GenericModel_coverage_obligations.json").read_text(encoding="utf-8"))
            self.assertTrue(checkpoint["evidence"]["probeExecuted"])
            self.assertEqual(after["summary"]["unresolved_count"], 0)

    def test_coverage_threshold_uses_all_three_metrics_for_every_model(self):
        report = {"M1": {"condition": {"percent": 90}, "decision": {"percent": 90}, "mcdc": {"percent": 79}}}
        self.assertFalse(RUNNER.coverage_meets(report, 80))
        report["M1"]["mcdc"]["percent"] = 80
        self.assertTrue(RUNNER.coverage_meets(report, 80))


if __name__ == "__main__":
    unittest.main()
