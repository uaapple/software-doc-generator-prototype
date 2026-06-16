#!/usr/bin/env python3
"""Backfill TCSD Action cells with simulation results for root outputs only."""

from __future__ import annotations

import argparse
import json
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


STEP_RE = re.compile(r"^\s*\[\+")
ANY_EXPECTED_RE = re.compile(r"^\s*([A-Za-z_]\w*)\s*=\s*expValue\(")
DEFAULT_UNSTABLE_POINT_DURATION_S = 0.01


def parse_steps(action: str) -> list[dict]:
    steps: list[dict] = []
    current: dict | None = None
    for raw in (action or "").splitlines():
        if STEP_RE.match(raw):
            if current is not None:
                steps.append(current)
            current = {"marker": raw, "lines": []}
        elif current is not None:
            current["lines"].append(raw)
    if current is not None:
        steps.append(current)
    for idx, step in enumerate(steps, start=1):
        step["index"] = idx
    return steps


def format_number(value: float) -> str:
    if abs(value) < 5e-8:
        return "0"
    rounded = round(value)
    if abs(value - rounded) < 5e-6:
        return str(int(rounded))
    return f"{value:.8g}"


def format_exp_value(value: float, duration: float | None = None, offset: float | None = None) -> str:
    formatted_value = format_number(value)
    if duration is None or offset is None:
        return f"expValue({formatted_value})"
    return f"expValue({formatted_value},{format_number(duration)},{format_number(offset)})"


def output_is_stable_for_step(step_result: dict, output: str) -> bool:
    """Return whether a single output can be trusted at this step."""
    return step_result.get("stable", {}).get(output) is not False


def build_action(
    action: str,
    step_results: dict[int, dict],
    outputs: list[str],
    unstable_point_duration_s: float = DEFAULT_UNSTABLE_POINT_DURATION_S,
    unstable_point_offset_s: float = 0.0,
) -> str:
    rebuilt: list[str] = []
    steps = parse_steps(action)
    for position, step in enumerate(steps):
        rebuilt.append(step["marker"])
        kept_lines: list[str] = []
        for line in step["lines"]:
            if ANY_EXPECTED_RE.match(line):
                continue
            kept_lines.append(line)
            rebuilt.append(line)
        is_final_empty_delay = position == len(steps) - 1 and not any(line.strip() for line in kept_lines)
        if is_final_empty_delay:
            continue
        result = step_results.get(step["index"], {})
        values = result.get("outputs", {})
        for output in outputs:
            if output not in values:
                continue
            value = float(values[output])
            if output_is_stable_for_step(result, output):
                rebuilt.append(f"{output} = {format_exp_value(value)};")
            elif unstable_point_duration_s > 0:
                rebuilt.append(
                    f"{output} = {format_exp_value(value, unstable_point_duration_s, unstable_point_offset_s)};"
                )
    return "\n".join(rebuilt)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--results", required=True)
    parser.add_argument("--outputs", required=True, help="Comma-separated root output names")
    parser.add_argument(
        "--exclude-outputs",
        default="",
        help="Comma-separated root outputs to remove from expValue backfill, for example unverified stateful outputs",
    )
    parser.add_argument(
        "--unstable-point-duration",
        type=float,
        default=DEFAULT_UNSTABLE_POINT_DURATION_S,
        help=(
            "Duration in seconds for expValue(value,duration,offset) when the full following interval is unstable. "
            "Use 0 to omit unstable-step point expectations."
        ),
    )
    parser.add_argument(
        "--unstable-point-offset",
        type=float,
        default=0.0,
        help="Offset in seconds for unstable-step expValue(value,duration,offset) point expectations.",
    )
    args = parser.parse_args()

    outputs = [name.strip() for name in args.outputs.split(",") if name.strip()]
    excluded_outputs = {name.strip() for name in args.exclude_outputs.split(",") if name.strip()}
    outputs = [name for name in outputs if name not in excluded_outputs]
    results = json.loads(Path(args.results).read_text(encoding="utf-8"))
    tests = results["tests"]
    if isinstance(tests, dict):
        tests = [tests]
    by_row = {}
    for item in tests:
        steps = item["steps"]
        if isinstance(steps, dict):
            steps = [steps]
        by_row[item["row"]] = {
            "steps": {step["index"]: step for step in steps},
        }

    wb = load_workbook(args.workbook)
    ws = wb["TCSD"]
    for row, info in by_row.items():
        cell = ws.cell(row, 7)
        cell.value = build_action(
            cell.value or "",
            info["steps"],
            outputs,
            unstable_point_duration_s=args.unstable_point_duration,
            unstable_point_offset_s=args.unstable_point_offset,
        )
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = "top"
        cell.alignment = alignment
        line_count = cell.value.count("\n") + 1
        ws.row_dimensions[row].height = min(409, max(180, line_count * 13))
    wb.save(args.workbook)
    print(args.workbook)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
