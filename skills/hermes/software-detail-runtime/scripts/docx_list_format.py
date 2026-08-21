#!/usr/bin/env python3
"""Normalize and validate native Word bullets in 实现方式 sections."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError as exc:  # pragma: no cover - environment error path
    raise SystemExit(
        "python-docx is required; use the bundled document workspace Python runtime"
    ) from exc


INTRO_BEFORE_PT = 7
INTRO_AFTER_PT = 2
BULLET_LEFT_DXA = 720
BULLET_HANGING_DXA = 360
BULLET_LINE_SPACING = 1.15
AUTO_LINE_SPACING_DXA = 276


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Normalize native level-zero bullets in DOCX 实现方式 sections."
    )
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--check-only", action="store_true")
    args = parser.parse_args()
    if not args.check_only and args.out is None:
        parser.error("--out is required unless --check-only is used")
    if args.check_only and args.out is not None:
        parser.error("--out cannot be combined with --check-only")
    return args


def heading_level(paragraph) -> int | None:
    name = paragraph.style.name.strip().lower()
    if not name.startswith("heading "):
        return None
    try:
        return int(name.split()[-1])
    except ValueError:
        return None


def implementation_sections(document) -> list[list]:
    sections: list[list] = []
    current: list | None = None
    for paragraph in document.paragraphs:
        level = heading_level(paragraph)
        if level is not None:
            if current is not None:
                sections.append(current)
                current = None
            if level == 3 and paragraph.text.strip() == "实现方式":
                current = []
            continue
        if current is not None:
            current.append(paragraph)
    if current is not None:
        sections.append(current)
    return sections


def is_literal_bullet(paragraph) -> bool:
    return paragraph.text.startswith("• ") or paragraph.text.startswith("•\t")


def is_native_bullet(paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def set_attr(element, name: str, value: str | int) -> None:
    element.set(qn(f"w:{name}"), str(value))


def create_bullet_numbering(document) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    set_attr(abstract, "abstractNumId", abstract_id)
    multi = OxmlElement("w:multiLevelType")
    set_attr(multi, "val", "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    set_attr(level, "ilvl", 0)
    for tag, value in (("w:start", 1), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
        node = OxmlElement(tag)
        set_attr(node, "val", value)
        level.append(node)
    level_ppr = OxmlElement("w:pPr")
    level_ind = OxmlElement("w:ind")
    set_attr(level_ind, "left", BULLET_LEFT_DXA)
    set_attr(level_ind, "hanging", BULLET_HANGING_DXA)
    level_ppr.append(level_ind)
    level.append(level_ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    set_attr(num, "numId", num_id)
    abstract_ref = OxmlElement("w:abstractNumId")
    set_attr(abstract_ref, "val", abstract_id)
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def strip_literal_marker(paragraph) -> None:
    remaining = 2
    for run in paragraph.runs:
        if remaining == 0:
            break
        if len(run.text) <= remaining:
            remaining -= len(run.text)
            run.text = ""
        else:
            run.text = run.text[remaining:]
            remaining = 0


def apply_native_bullet(paragraph, num_id: int | None = None) -> None:
    if is_literal_bullet(paragraph):
        strip_literal_marker(paragraph)
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.get_or_add_numPr()
    ilvl = num_pr.get_or_add_ilvl()
    ilvl.val = 0
    if num_id is not None:
        num_pr.get_or_add_numId().val = num_id
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = BULLET_LINE_SPACING


def normalize(document) -> dict[str, int]:
    sections = implementation_sections(document)
    literal_count = sum(is_literal_bullet(p) for section in sections for p in section)
    num_id = create_bullet_numbering(document) if literal_count else None

    removed_empty = 0
    for section in sections:
        for paragraph in section:
            if not paragraph.text.strip():
                remove_paragraph(paragraph)
                removed_empty += 1

    native_count = 0
    for section in implementation_sections(document):
        for paragraph in section:
            if is_literal_bullet(paragraph):
                apply_native_bullet(paragraph, num_id)
                native_count += 1
            elif is_native_bullet(paragraph):
                apply_native_bullet(paragraph)
                native_count += 1

        refreshed = [p for p in section if p._element.getparent() is not None]
        for index in range(len(refreshed) - 1):
            current = refreshed[index]
            following = refreshed[index + 1]
            if current.text.strip() and not is_native_bullet(current) and is_native_bullet(following):
                current.paragraph_format.space_before = Pt(INTRO_BEFORE_PT)
                current.paragraph_format.space_after = Pt(INTRO_AFTER_PT)
                current.paragraph_format.keep_with_next = True

    return {
        "sections": len(sections),
        "native_bullets": native_count,
        "removed_empty_paragraphs": removed_empty,
    }


def direct_twips(element, name: str) -> int | None:
    if element is None:
        return None
    value = element.get(qn(f"w:{name}"))
    return int(value) if value is not None else None


def validate(document) -> list[str]:
    errors: list[str] = []
    sections = implementation_sections(document)
    bullet_total = 0
    for section_index, section in enumerate(sections, start=1):
        groups: list[int] = []
        group_size = 0
        previous = None
        for paragraph_index, paragraph in enumerate(section, start=1):
            label = f"section {section_index}, paragraph {paragraph_index}"
            if not paragraph.text.strip():
                errors.append(f"{label}: empty paragraphs are forbidden in 实现方式")
            if is_literal_bullet(paragraph):
                errors.append(f"{label}: literal bullet prefix is forbidden")
            native = is_native_bullet(paragraph)
            if native:
                bullet_total += 1
                group_size += 1
                ppr = paragraph._p.pPr
                ilvl = ppr.numPr.ilvl
                if ilvl is None or int(ilvl.val) != 0:
                    errors.append(f"{label}: native bullet must use level 0")
                if ppr.numPr.numId is None:
                    errors.append(f"{label}: native bullet must have a numbering ID")
                if direct_twips(ppr.ind, "left") != BULLET_LEFT_DXA:
                    errors.append(f"{label}: left indent must be {BULLET_LEFT_DXA} DXA")
                if direct_twips(ppr.ind, "hanging") != BULLET_HANGING_DXA:
                    errors.append(f"{label}: hanging indent must be {BULLET_HANGING_DXA} DXA")
                spacing = ppr.spacing
                if direct_twips(spacing, "before") != 0:
                    errors.append(f"{label}: bullet space before must be 0 pt")
                if direct_twips(spacing, "after") != 0:
                    errors.append(f"{label}: bullet space after must be 0 pt")
                if direct_twips(spacing, "line") != AUTO_LINE_SPACING_DXA:
                    errors.append(f"{label}: bullet line spacing must be 1.15")
                if spacing is None or spacing.get(qn("w:lineRule")) != "auto":
                    errors.append(f"{label}: bullet line spacing rule must be auto")
                if previous is not None and not previous.text.strip():
                    errors.append(f"{label}: empty paragraph before bullet")
            elif group_size:
                groups.append(group_size)
                group_size = 0
            if (
                not native
                and paragraph.text.strip()
                and paragraph_index < len(section)
                and is_native_bullet(section[paragraph_index])
            ):
                ppr = paragraph._p.pPr
                spacing = ppr.spacing if ppr is not None else None
                if direct_twips(spacing, "before") != INTRO_BEFORE_PT * 20:
                    errors.append(f"{label}: introduction space before must be {INTRO_BEFORE_PT} pt")
                if direct_twips(spacing, "after") != INTRO_AFTER_PT * 20:
                    errors.append(f"{label}: introduction space after must be {INTRO_AFTER_PT} pt")
                if ppr is None or ppr.keepNext is None:
                    errors.append(f"{label}: introduction must keep with the first bullet")
            previous = paragraph
        if group_size:
            groups.append(group_size)
        for group_index, size in enumerate(groups, start=1):
            if not 2 <= size <= 5:
                errors.append(
                    f"section {section_index}, group {group_index}: expected 2-5 bullets, found {size}"
                )
    if not sections:
        errors.append("no 实现方式 section found")
    if bullet_total == 0:
        errors.append("no native bullets found in 实现方式 sections")
    return errors


def main() -> int:
    args = parse_args()
    if not args.input.is_file():
        print(f"ERROR: DOCX not found: {args.input}", file=sys.stderr)
        return 2
    document = Document(args.input)
    stats = None
    if not args.check_only:
        stats = normalize(document)
        args.out.parent.mkdir(parents=True, exist_ok=True)
        document.save(args.out)
        document = Document(args.out)
    errors = validate(document)
    if errors:
        for error in errors:
            print(f"ERROR: {error}", file=sys.stderr)
        return 1
    if stats is None:
        stats = {
            "sections": len(implementation_sections(document)),
            "native_bullets": sum(
                is_native_bullet(p)
                for section in implementation_sections(document)
                for p in section
            ),
            "removed_empty_paragraphs": 0,
        }
    print(
        "PASS: {sections} implementation section(s); {native_bullets} native bullet(s); "
        "{removed_empty_paragraphs} empty paragraph(s) removed.".format(**stats)
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
