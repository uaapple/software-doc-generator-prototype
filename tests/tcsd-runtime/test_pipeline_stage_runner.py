import importlib.util
import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

RUNTIME = Path(__file__).resolve().parents[2] / "skills" / "hermes" / "tcsd-runtime"
SCRIPT = RUNTIME / "scripts" / "run_tcsd_pipeline_stage.py"
SATK_SCRIPT = RUNTIME / "scripts" / "satk_eval.py"
REPAIR_SCRIPT = RUNTIME / "scripts" / "validate_agent_coverage_repair.py"
SPEC = importlib.util.spec_from_file_location("run_tcsd_pipeline_stage", SCRIPT)
RUNNER = importlib.util.module_from_spec(SPEC)
assert SPEC.loader
SPEC.loader.exec_module(RUNNER)
SATK_SPEC = importlib.util.spec_from_file_location("satk_eval", SATK_SCRIPT)
SATK = importlib.util.module_from_spec(SATK_SPEC)
assert SATK_SPEC.loader
SATK_SPEC.loader.exec_module(SATK)
REPAIR_SPEC = importlib.util.spec_from_file_location("validate_agent_coverage_repair", REPAIR_SCRIPT)
REPAIR = importlib.util.module_from_spec(REPAIR_SPEC)
assert REPAIR_SPEC.loader
REPAIR_SPEC.loader.exec_module(REPAIR)


class PipelineStageRunnerTests(unittest.TestCase):
    def test_stage10_agent_repair_preserves_focused_temporal_stimulus(self):
        brief = {
            "schema": REPAIR.BRIEF_SCHEMA,
            "jobId": "job-generic",
            "model": "GenericModel",
            "repairRequired": True,
            "metricDeficits": [{"coverage_class": "Decision"}],
            "coverageTargets": [{
                "coverage_class": "Decision",
                "block": {"path": "GenericModel/Decision", "sid": "GenericModel:7"},
                "missing_outcomes": ["false branch after timeout"],
                "requires_model_inspection": False,
            }],
            "guardrails": {
                "maxCandidateTests": 16,
                "maxStepsPerTest": 8,
                "parametersOnlyInInitialization": True,
                "analyzeOnlyTargetUpstreamSlice": True,
                "fullRootInputEnumerationForbidden": True,
                "repairPassLimit": 1,
            },
        }
        proposal = {
            "schema": REPAIR.PROPOSAL_SCHEMA,
            "jobId": "job-generic",
            "model": "GenericModel",
            "tests": [{
                "id": "decision-timeout-false",
                "coverage_class": "Decision",
                "block": {"path": "GenericModel/Decision", "sid": "GenericModel:7"},
                "required_outcome": "false branch after timeout",
                "controller": {
                    "direct_inputs": {"Enable": 1},
                    "parameters": {"WaitThreshold": 3},
                },
                "stimulus": {
                    "initial_inputs": {"Enable": 0},
                    "initial_params": {"Bypass": 0},
                    "steps": [
                        {"delay_s": 0.1, "input_updates": {"Enable": 1}, "param_updates": {}},
                        {"delay_s": 0.3, "input_updates": {}, "param_updates": {}},
                    ],
                    "evidence_step": 2,
                },
                "analysis": {
                    "upstream_slice": ["GenericModel/Delay", "GenericModel/Decision"],
                    "rationale": "Enable is asserted, held across the threshold, and unrelated gates stay sensitized.",
                },
            }],
            "unresolved": [],
        }
        ir, report = REPAIR.validate_proposal(
            proposal,
            brief,
            {"schema": "tcsd-model-interface/v1", "inputs": ["Enable"], "outputs": ["Output"]},
        )
        self.assertEqual(report["acceptedCandidateCount"], 1)
        self.assertEqual(ir["items"][0]["stimulus"]["evidence_step"], 2)
        self.assertEqual(ir["items"][0]["analysis"]["cumulative_wait_s"], 0.4)
        self.assertEqual(ir["items"][0]["controller"]["parameters"], {"Bypass": 0, "WaitThreshold": 3})

        proposal["tests"][0]["required_outcome"] = "unmeasured true branch"
        with self.assertRaisesRegex(ValueError, "measured missing outcome"):
            REPAIR.validate_proposal(
                proposal,
                brief,
                {"schema": "tcsd-model-interface/v1", "inputs": ["Enable"], "outputs": ["Output"]},
            )
        proposal["tests"][0]["required_outcome"] = "false branch after timeout"
        proposal["tests"][0]["stimulus"]["steps"][1]["param_updates"] = {"Bypass": 1}
        with self.assertRaisesRegex(ValueError, "only be set in initialization"):
            REPAIR.validate_proposal(
                proposal,
                brief,
                {"schema": "tcsd-model-interface/v1", "inputs": ["Enable"], "outputs": ["Output"]},
            )

    def test_stage10_unresolved_requires_specific_reason_and_evidence(self):
        brief = {
            "schema": REPAIR.BRIEF_SCHEMA,
            "jobId": "job-generic",
            "model": "GenericModel",
            "repairRequired": True,
            "metricDeficits": [{"coverage_class": "Decision"}],
            "coverageTargets": [],
            "guardrails": {
                "maxCandidateTests": 16,
                "maxStepsPerTest": 8,
                "stepCountSemantics": "stimulus_action_entries",
                "simulationSamplePeriodsDoNotCountAsSteps": True,
                "longHoldAsSingleActionAllowed": True,
            },
        }
        proposal = {
            "schema": REPAIR.PROPOSAL_SCHEMA,
            "jobId": "job-generic",
            "model": "GenericModel",
            "tests": [],
            "unresolved": [{
                "coverage_class": "Decision",
                "block": {"path": "GenericModel/Decision", "sid": "GenericModel:7"},
                "reason_code": "state_sequence_not_constructible",
                "evidence": "The reset prerequisite is controlled by an unmodifiable literal.",
            }],
        }
        ir, report = REPAIR.validate_proposal(
            proposal,
            brief,
            {"schema": "tcsd-model-interface/v1", "inputs": ["Enable"], "outputs": ["Output"]},
        )
        self.assertEqual(ir["summary"]["unresolved"], 1)
        self.assertEqual(report["unresolved"][0]["reason_code"], "state_sequence_not_constructible")
        proposal["unresolved"][0]["reason_code"] = "no_candidate"
        with self.assertRaisesRegex(ValueError, "specific allowed reason"):
            REPAIR.validate_proposal(
                proposal,
                brief,
                {"schema": "tcsd-model-interface/v1", "inputs": ["Enable"], "outputs": ["Output"]},
            )

    def test_stage10_rejects_sample_period_count_as_action_step_limit(self):
        brief = {
            "schema": REPAIR.BRIEF_SCHEMA,
            "jobId": "job-counter",
            "model": "GenericCounter",
            "repairRequired": True,
            "metricDeficits": [{"coverage_class": "Decision"}],
            "coverageTargets": [{
                "coverage_class": "Decision",
                "block": {"path": "GenericCounter/MinMax", "sid": "GenericCounter:15"},
                "missing_outcomes": ["input 2"],
                "requires_model_inspection": False,
            }],
            "guardrails": {
                "maxCandidateTests": 16,
                "maxStepsPerTest": 8,
                "stepCountSemantics": "stimulus_action_entries",
                "simulationSamplePeriodsDoNotCountAsSteps": True,
                "longHoldAsSingleActionAllowed": True,
            },
        }
        proposal = {
            "schema": REPAIR.PROPOSAL_SCHEMA,
            "jobId": "job-counter",
            "model": "GenericCounter",
            "tests": [],
            "unresolved": [{
                "coverage_class": "Decision",
                "block": {"path": "GenericCounter/MinMax", "sid": "GenericCounter:15"},
                "reason_code": "state_sequence_not_constructible",
                "evidence": (
                    "The counter needs 65535 simulation sample steps, far exceeding "
                    "the 8-step per-test guardrail."
                ),
            }],
        }
        with self.assertRaisesRegex(ValueError, "sample periods as TCSD action steps"):
            REPAIR.validate_proposal(
                proposal,
                brief,
                {"schema": "tcsd-model-interface/v1", "inputs": ["Enable"], "outputs": ["Output"]},
            )

    def test_stage10_accepts_long_hold_as_one_action_step(self):
        brief = {
            "schema": REPAIR.BRIEF_SCHEMA,
            "jobId": "job-counter",
            "model": "GenericCounter",
            "repairRequired": True,
            "metricDeficits": [{"coverage_class": "Decision"}],
            "coverageTargets": [{
                "coverage_class": "Decision",
                "block": {"path": "GenericCounter/MinMax", "sid": "GenericCounter:15"},
                "missing_outcomes": ["input 2"],
                "requires_model_inspection": False,
            }],
            "guardrails": {
                "maxCandidateTests": 16,
                "maxStepsPerTest": 8,
                "stepCountSemantics": "stimulus_action_entries",
                "simulationSamplePeriodsDoNotCountAsSteps": True,
                "longHoldAsSingleActionAllowed": True,
            },
        }
        proposal = {
            "schema": REPAIR.PROPOSAL_SCHEMA,
            "jobId": "job-counter",
            "model": "GenericCounter",
            "tests": [{
                "id": "counter-input2",
                "coverage_class": "Decision",
                "block": {"path": "GenericCounter/MinMax", "sid": "GenericCounter:15"},
                "required_outcome": "input 2",
                "controller": {"direct_inputs": {"Enable": 1}, "parameters": {}},
                "stimulus": {
                    "initial_inputs": {"Enable": 1},
                    "initial_params": {},
                    "steps": [{"delay_s": 655.36, "input_updates": {}, "param_updates": {}}],
                    "evidence_step": 1,
                },
                "analysis": {
                    "upstream_slice": [
                        "GenericCounter/MinMax",
                        "GenericCounter/Sum",
                        "GenericCounter/Unit Delay",
                    ],
                    "rationale": (
                        "A 0.01 second sample time needs 65536 sample periods, "
                        "represented by one finite hold action."
                    ),
                },
            }],
            "unresolved": [],
        }
        ir, report = REPAIR.validate_proposal(
            proposal,
            brief,
            {"schema": "tcsd-model-interface/v1", "inputs": ["Enable"], "outputs": ["Output"]},
        )
        self.assertEqual(report["acceptedCandidateCount"], 1)
        self.assertEqual(len(ir["items"][0]["stimulus"]["steps"]), 1)
        self.assertEqual(ir["items"][0]["analysis"]["cumulative_wait_s"], 655.36)

    def test_satk_server_discovery_priority_is_cross_platform_and_deterministic(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            home = root / "home"
            toolkit_bin = home / ".matlab" / "agentic-toolkits" / "bin"
            repository = root / "repository"
            script_file = repository / "skills" / "hermes" / "tcsd-runtime" / "scripts" / "satk_eval.py"
            toolkit_bin.mkdir(parents=True)
            (repository / "tools").mkdir(parents=True)
            official = toolkit_bin / "matlab-mcp-server"
            legacy = toolkit_bin / "matlab-mcp-core-server"
            tools_legacy = repository / "tools" / "matlab-mcp-core-server"
            explicit = root / "explicit-mcp"
            for candidate in (official, legacy, tools_legacy, explicit):
                candidate.write_bytes(candidate.name.encode("utf-8"))

            selected, source = SATK.resolve_server(
                environ={"SATK_MCP_SERVER": str(explicit)},
                home=home,
                cwd=root,
                script_file=script_file,
                platform_name="Darwin",
            )
            self.assertEqual((selected, source), (explicit.resolve(), "environment"))

            selected, source = SATK.resolve_server(
                environ={},
                home=home,
                cwd=root,
                script_file=script_file,
                platform_name="Darwin",
            )
            self.assertEqual((selected, source), (official.resolve(), "official-toolkit"))
            official.unlink()
            selected, source = SATK.resolve_server(
                environ={},
                home=home,
                cwd=root,
                script_file=script_file,
                platform_name="Darwin",
            )
            self.assertEqual((selected, source), (legacy.resolve(), "legacy-toolkit"))
            legacy.unlink()
            selected, source = SATK.resolve_server(
                environ={},
                home=home,
                cwd=root,
                script_file=script_file,
                platform_name="Darwin",
            )
            self.assertEqual((selected, source), (tools_legacy.resolve(), "repository-tools-legacy"))
            self.assertEqual(
                SATK.server_binary_names("Windows"),
                ("matlab-mcp-server.exe", "matlab-mcp-core-server.exe"),
            )

    def test_satk_runner_treats_tool_result_is_error_as_failure(self):
        self.assertTrue(SATK.mcp_response_failed({"error": {"code": -1}}))
        self.assertTrue(SATK.mcp_response_failed({"result": {"isError": True, "content": []}}))
        self.assertFalse(SATK.mcp_response_failed({"result": {"isError": False, "content": []}}))

    def test_stage02_extracts_public_mcp_error_from_satk_failure(self):
        response = json.dumps({
            "jsonrpc": "2.0",
            "id": 2,
            "result": {
                "isError": True,
                "content": [{
                    "type": "text",
                    "text": "failed to attach to MATLAB session",
                }],
            },
        })
        failure = subprocess.CalledProcessError(
            1,
            ["python3", "satk_eval.py", "canary.m"],
            output=response,
            stderr="",
        )
        with mock.patch.object(RUNNER.subprocess, "run", side_effect=failure):
            with self.assertRaisesRegex(
                RuntimeError,
                r"Stage 02 environment gate failed: SATK/MCP failed: failed to attach to MATLAB session",
            ):
                RUNNER.run_satk(
                    ["python3", "satk_eval.py", "canary.m"],
                    Path("/tmp"),
                    stage=2,
                    context="environment gate failed",
                )

    def test_stage02_matlab_root_falls_back_to_satk_root(self):
        with mock.patch.dict(
            os.environ,
            {"SATK_MATLAB_ROOT": "/Applications/MATLAB_R2026a.app"},
            clear=True,
        ):
            self.assertEqual(
                RUNNER.matlab_root_path({"matlabRoot": "/ignored/input/root"}),
                Path("/Applications/MATLAB_R2026a.app"),
            )

    def test_public_stage_error_redacts_credentials(self):
        self.assertEqual(
            RUNNER.public_error_text("failed password=visible token: also-visible"),
            "failed password=[REDACTED] token: [REDACTED]",
        )

    def test_planning_mapping_gaps_are_advisory_and_superseded_by_measured_coverage(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            obligations = root / "outputs" / "GenericModel_planning_obligations_snapshot.json"
            obligations.parent.mkdir()
            obligations.write_text('{"obligations":[]}', encoding="utf-8")
            assessment = RUNNER.planning_mapping_assessment(
                {
                    "status": "failed",
                    "summary": {"missing_count": 3, "unresolved_count": 1},
                    "missing": [{"id": "generic-obligation"}],
                },
                obligations,
                root,
            )
            self.assertEqual(assessment["schema"], "tcsd-planning-mapping-assessment/v1")
            self.assertEqual(assessment["status"], "advisory")
            self.assertEqual(assessment["assessment"], "gaps-observed")
            self.assertFalse(assessment["blocking"])
            self.assertNotEqual(assessment["status"], "failed")
            self.assertEqual(assessment["supersededBy"]["stageIndex"], 9)
            self.assertEqual(assessment["supersededBy"]["authority"], "measured-simulink-coverage")
            self.assertEqual(assessment["sourceObligations"]["sha256"], hashlib.sha256(obligations.read_bytes()).hexdigest())

    def test_first_three_stages_write_candidate_results_but_not_host_checkpoints(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            output = root / "outputs"
            model = root / "GenericModel.slx"
            mat = root / "GenericModel.mat"
            model.write_bytes(b"slx")
            mat.write_bytes(b"mat")
            job = {
                "jobId": "job-generic",
                "events": [],
                "resources": {"ownerJobId": "job-generic"},
                "input": {
                    "workspaceDir": str(root),
                    "outputDir": str(output),
                    "modelSlxPath": str(model),
                    "modelMatPath": str(mat),
                    "projectAddonCopy": {"copiedFileCount": 0},
                },
            }
            fake_matlab_root = root / "fake-matlab"
            (fake_matlab_root / "bin").mkdir(parents=True)
            (fake_matlab_root / "bin" / "matlab").write_text("fixture", encoding="utf-8")
            fake_mcp_server = root / "fake-matlab-mcp-server"
            fake_mcp_server.write_bytes(b"offline MCP fixture")
            job["input"]["matlabRoot"] = str(fake_matlab_root)
            environment_fixture = root / "environment-fixture.json"
            nonce = "offline-fixture-nonce"
            environment_fixture.write_text(json.dumps({
                "schema": "tcsd-environment-gate/v2",
                "jobId": "job-generic",
                "nonce": nonce,
                "matlabRoot": str(fake_matlab_root),
                "python": sys.executable,
                "runner": str(SATK_SCRIPT),
                "pythonDependencies": {
                    "passed": True,
                    "modules": {
                        "yaml": {"version": "6.0.3"},
                        "openpyxl": {"version": "3.1.5"},
                    },
                },
                "workspaceIo": {"passed": True, "created": True, "readMatched": True, "deleted": True},
                "matlab": {"passed": True, "nonce": nonce, "version": "R2026a"},
                "simulink": {
                    "passed": True,
                    "licenseAvailable": True,
                    "loaded": True,
                    "version": "R2026a",
                },
                "satkMcp": {
                    "passed": True,
                    "runner": str(SATK_SCRIPT),
                    "server": {
                        "path": str(fake_mcp_server),
                        "discovery": "offline-test-fixture",
                        "sha256": hashlib.sha256(fake_mcp_server.read_bytes()).hexdigest(),
                        "sizeBytes": fake_mcp_server.stat().st_size,
                    },
                    "sentinelWritten": True,
                    "nonceMatched": True,
                },
                "passed": True,
            }), encoding="utf-8")
            old = os.environ.get("TCSD_PIPELINE_ENV_CANARY_FIXTURE")
            old_setup = os.environ.get("TCSD_PIPELINE_SETUP_FIXTURE")
            os.environ["TCSD_PIPELINE_ENV_CANARY_FIXTURE"] = str(environment_fixture)
            os.environ["TCSD_PIPELINE_SETUP_FIXTURE"] = "1"
            try:
                for stage in (1, 2, 3):
                    result_path = output / ".tcsd-results" / f"stage-{stage:02d}.json"
                    job["_stageResultPath"] = str(result_path)
                    RUNNER.stage_run(stage, job)
                    result = json.loads(result_path.read_text(encoding="utf-8"))
                    self.assertEqual(result["schema"], "tcsd-agent-stage-result/v1")
                    self.assertEqual(result["jobId"], "job-generic")
                    self.assertEqual(result["stageIndex"], stage)
                    self.assertTrue(result["artifacts"])
                    self.assertFalse((output / ".tcsd-checkpoints").exists())
                resources = json.loads((output / ".tcsd-runtime" / "owned-resources.json").read_text(encoding="utf-8"))
                self.assertEqual(resources["jobId"], "job-generic")
                initialized = json.loads((output / ".tcsd-evidence" / "workspace-initialization.json").read_text(encoding="utf-8"))
                self.assertEqual(initialized["schema"], "tcsd-workspace-initialization/v1")
                self.assertTrue(initialized["completed"])
            finally:
                if old is None:
                    os.environ.pop("TCSD_PIPELINE_ENV_CANARY_FIXTURE", None)
                else:
                    os.environ["TCSD_PIPELINE_ENV_CANARY_FIXTURE"] = old
                if old_setup is None:
                    os.environ.pop("TCSD_PIPELINE_SETUP_FIXTURE", None)
                else:
                    os.environ["TCSD_PIPELINE_SETUP_FIXTURE"] = old_setup

    def test_single_input_output_interface_is_not_split_into_characters(self):
        interface = RUNNER.validate_interface({"schema": "tcsd-model-interface/v1", "inputs": ["OnlyInput"], "outputs": ["OnlyOutput"]})
        spec = RUNNER.initial_spec({"inputs": "OnlyInput", "outputs": "OnlyOutput"}, "GenericModel")
        self.assertEqual(interface["inputs"], ["OnlyInput"])
        self.assertEqual(spec["tests"][0]["initialization"], "OnlyInput=0;")
        self.assertNotIn("output_reference", spec["tests"][0])
        with self.assertRaises(RuntimeError):
            RUNNER.validate_interface({"schema": "tcsd-model-interface/v1", "inputs": "OnlyInput", "outputs": ["OnlyOutput"]})

    def test_stage_four_initializes_its_new_session_before_tracing(self):
        code = RUNNER.stage4_matlab_code(root=Path("C:/job"), scripts_dir=Path("C:/skill/scripts"), interface=Path("C:/job/outputs/interface.json"), model="GenericModel", mat_name="GenericModel.mat", init_scripts=["project_init.m"])
        setup = code.index("setup_ut_support(rootDir,initScripts)")
        trace = code.index("trace_logical_mcdc")
        self.assertIn("initScripts={'project_init.m'}", code)
        self.assertIn("'WorkspaceInitialized',true", code)
        self.assertLess(setup, trace)

    def test_simulation_backfill_requires_matching_real_result_counts(self):
        with tempfile.TemporaryDirectory() as temp:
            workbook = Path(temp) / "result.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "TCSD"; ws["A3"] = "TC_001"; ws["C3"] = "Test"; ws["G3"] = "[+0.1s]\nOnlyOutput = expValue(1);\n[+0.1s]"; wb.save(workbook)
            simulation = {"tests": [{"row": 3, "test_id": "TC_001", "steps": [{"index": 1, "outputs": {"OnlyOutput": 1}, "stable": {"OnlyOutput": True}}, {"index": 2, "outputs": {"OnlyOutput": 1}, "stable": {"OnlyOutput": True}}]}]}
            evidence = RUNNER.simulation_backfill_evidence(simulation, workbook)
            self.assertEqual(evidence["simulationValueCount"], 1)
            self.assertEqual(evidence["workbookBackfillCount"], 1)
            self.assertEqual(evidence["backfillItems"], [{"row": 3, "testId": "TC_001", "step": 1, "output": "OnlyOutput", "value": 1.0}])
            with self.assertRaises(RuntimeError):
                RUNNER.simulation_backfill_evidence({"tests": []}, workbook)

    def test_simulation_backfill_rejects_missing_extra_and_wrong_values(self):
        with tempfile.TemporaryDirectory() as temp:
            workbook = Path(temp) / "result.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "TCSD"; ws["A4"] = "TC_002"; ws["C4"] = "Test"
            simulation = {"tests": [{"row": 4, "test_id": "TC_002", "steps": [{"index": 1, "outputs": {"Result": 2.5}, "stable": {"Result": True}}]}]}
            for action in ("[+0.1s]", "[+0.1s]\nResult = expValue(2.5);\nExtra = expValue(1);", "[+0.1s]\nResult = expValue(2.4);"):
                ws["G4"] = action; wb.save(workbook)
                with self.assertRaises(RuntimeError): RUNNER.simulation_backfill_evidence(simulation, workbook)

    def test_stage_five_to_six_cli_persists_unresolved_then_runs_probe(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); output = root / "outputs"; output.mkdir()
            model = root / "GenericModel.slx"; mat = root / "GenericModel.mat"; model.write_bytes(b"slx"); mat.write_bytes(b"mat")
            trace = {
                "schema": "simulink-ut-logical-mcdc-trace/v2", "model": "GenericModel", "operator_count": 1,
                "operators": [{"id": "GenericModel:1", "sid": "GenericModel:1", "block_path": "GenericModel/Decision", "operator": "AND", "ports": [
                    {"index": 1, "trace": {"kind": "relational", "operator": ">", "inputs": [{"trace": {"kind": "stateful", "block_type": "UnitDelay", "inputs": [{"trace": {"kind": "root_inport", "signal": "Enable"}}]}}, {"trace": {"kind": "constant", "value": "0"}}]}},
                    {"index": 2, "trace": {"kind": "root_inport", "signal": "Request"}},
                ]}],
            }
            (output / "GenericModel_logical_traces.json").write_text(json.dumps(trace), encoding="utf-8")
            probe = {"GenericModel": {"schema": "simulink-ut-logical-mcdc-probe/v2", "model": "GenericModel", "probes": [{"id": "GenericModel:1", "sid": "GenericModel:1", "block_path": "GenericModel/Decision", "operator": "AND", "port_names": ["u1", "u2"]}], "observations": []}}
            for index, label in enumerate(("TT", "FT", "TF"), 1):
                probe["GenericModel"]["observations"].append({"test_id": f"STATE_PROBE_{index:04d}", "row": index, "step_index": 2, "time_s": 0.1, "inputs": {"Enable": int(label[0] == "T"), "Request": int(label[1] == "T")}, "params": {}, "vectors": {"decision": {"id": "GenericModel:1", "label": label, "ok": True}}, "stimulus": {"initial_inputs": {"Enable": 0, "Request": 1}, "initial_params": {}, "steps": [{"index": 1, "delay_s": 0.01, "input_updates": {"Enable": 1}, "param_updates": {}}, {"index": 2, "delay_s": 0.1, "input_updates": {}, "param_updates": {}}], "evidence_step": 2}, "prediction_status": "observed"})
            probe_fixture = root / "probe-results.json"; probe_fixture.write_text(json.dumps(probe), encoding="utf-8")
            job = {"jobId": "job-cli", "resources": {"ownerJobId": "job-cli"}, "input": {"workspaceDir": str(root), "outputDir": str(output), "modelSlxPath": str(model), "modelMatPath": str(mat), "coverageThreshold": 80}}
            def run_stage(stage, env=None):
                manifest_path = root / f"stage-{stage:02d}-input.json"
                result_path = output / ".tcsd-results" / f"stage-{stage:02d}.json"
                manifest_path.write_text(
                    json.dumps(
                        {
                            "schema": "tcsd-agent-stage-input/v1",
                            "jobId": "job-cli",
                            "stageIndex": stage,
                            "attempt": 1,
                            "job": job,
                        }
                    ),
                    encoding="utf-8",
                )
                subprocess.run(
                    [sys.executable, str(SCRIPT), "--manifest", str(manifest_path), "--result", str(result_path)],
                    check=True,
                    cwd=root,
                    env=env,
                )
                return json.loads(result_path.read_text(encoding="utf-8"))

            run_stage(5)
            before = json.loads((output / "GenericModel_coverage_obligations.json").read_text(encoding="utf-8"))
            self.assertGreater(before["summary"]["unresolved_count"], 0)
            env = dict(os.environ); env["TCSD_PIPELINE_PROBE_RESULTS_FIXTURE"] = str(probe_fixture)
            result = run_stage(6, env=env)
            after = json.loads((output / "GenericModel_coverage_obligations.json").read_text(encoding="utf-8"))
            self.assertEqual(result["schema"], "tcsd-agent-stage-result/v1")
            self.assertTrue(result["evidence"]["probeExecuted"])
            self.assertFalse((output / ".tcsd-checkpoints").exists())
            self.assertEqual(after["summary"]["unresolved_count"], 0)

    def test_coverage_threshold_uses_all_three_metrics_for_every_model(self):
        report = {"M1": {"condition": {"percent": 90}, "decision": {"percent": 90}, "mcdc": {"percent": 79}}}
        self.assertFalse(RUNNER.coverage_meets(report, 80))
        report["M1"]["mcdc"]["percent"] = 80
        self.assertTrue(RUNNER.coverage_meets(report, 80))


class StageRunnerHostFixesTests(unittest.TestCase):
    """Regression tests for the host-side runner fixes discovered on the
    RngPrdn_A02_B04 real-model run (semantic interface fallback and finish
    manifests with real coverage/repair facts)."""

    @classmethod
    def setUpClass(cls) -> None:
        spec = importlib.util.spec_from_file_location(
            "dsh_stage_runner", RUNTIME / "scripts" / "dsh_stage_runner.py")
        cls.runner = importlib.util.module_from_spec(spec)
        assert spec.loader
        spec.loader.exec_module(cls.runner)

    def test_semantic_validate_falls_back_to_outputs_interface(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            output = root / "outputs"
            output.mkdir()
            interface = output / "GenericModel_interface.json"
            interface.write_text(json.dumps({"schema": "tcsd-model-interface/v1"}), encoding="utf-8")
            task = {"id": "job-if", "workspace": {"directory": str(root)}}
            result = {
                "schema": "tcsd-agent-stage-result/v1",
                "jobId": "job-if",
                "stageIndex": 7,
                "status": "completed",
                "artifacts": [{"path": "outputs/GenericModel_tcsd_spec.json", "kind": "json", "role": "output"}],
            }
            request_path = output / "semantic-request.json"
            # Fake the host validator subprocess: it must receive the fallback interface path.
            captured = {}

            def fake_run(cmd, cwd, env, capture_output, text):
                req = json.loads(request_path.read_text(encoding="utf-8"))
                captured["interfacePath"] = req.get("interfacePath")
                return mock.Mock(returncode=0, stdout=json.dumps({
                    "schema": "tcsd-host-semantic-validation/v1", "stageIndex": 7,
                    "passed": True, "details": {}}), stderr="")

            with mock.patch.object(self.runner.subprocess, "run", side_effect=fake_run):
                report = self.runner.semantic_validate(
                    task, 7, result, RUNTIME, request_path, output / "semantic-validation.json")
            self.assertTrue(report["passed"])
            self.assertEqual(captured["interfacePath"], str(interface))

    def test_finish_writes_real_coverage_repair_and_picks_highest_iter(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workspace = root / "workspace"
            output = workspace / "outputs"
            model_dir = root / "model-dir"
            (output / ".tcsd-host").mkdir(parents=True)
            model_dir.mkdir()
            model = "RngPrdn_A02_B04"
            iter0 = output / f"{model}_Test_coverage_ir_iter0.xlsx"
            iter1 = output / f"{model}_Test_coverage_ir_iter1.xlsx"
            plain = output / f"{model}_Test0001_tcsd.xlsx"
            for path in (iter0, iter1, plain):
                wb = Workbook()
                wb.save(path)
            initial = {"models": {model: {"condition": {"covered": 15, "total": 30, "percent": 50, "passed": False},
                                          "decision": {"covered": 11, "total": 20, "percent": 55, "passed": False},
                                          "mcdc": {"covered": 0, "total": 6, "percent": 0, "passed": False},
                                          "test_count": 4, "threshold": 80}}}
            final = {"models": {model: {"condition": {"covered": 29, "total": 30, "percent": 96.67, "passed": True},
                                         "decision": {"covered": 20, "total": 20, "percent": 100, "passed": True},
                                         "mcdc": {"covered": 4, "total": 6, "percent": 66.67, "passed": False},
                                         "test_count": 13, "threshold": 80}}}
            (output / f"{model}_initial_coverage_summary.json").write_text(json.dumps(initial), encoding="utf-8")
            (output / f"{model}_final_coverage_summary.json").write_text(json.dumps(final), encoding="utf-8")
            (output / f"{model}_repair_candidate_validation.json").write_text(json.dumps(
                {"schema": "tcsd-repair-candidate-validation/v1", "jobId": "job-x",
                 "passed": True, "candidateCount": 9}), encoding="utf-8")
            (output / f"{model}_agent_coverage_repair_proposal.json").write_text(json.dumps(
                {"schema": "tcsd-agent-coverage-repair-proposal/v1", "jobId": "job-x", "model": model,
                 "tests": [], "unresolved": [
                     {"coverage_class": "MCDC",
                      "block": {"path": f"{model}/RampLimiter2/Logical Operator1", "sid": "231"},
                      "reason_code": "logic_unreachable",
                      "evidence": "algebraic: 233 true implies 234 true"}]}), encoding="utf-8")
            checkpoints = output / ".tcsd-checkpoints"
            checkpoints.mkdir()
            for stage in range(1, 13):
                (checkpoints / f"stage-{stage:02d}.json").write_text(json.dumps(
                    {"stageIndex": stage, "status": "completed", "attempt": 1, "summary": "s"}), encoding="utf-8")
            task = {"id": "job-x",
                    "workspace": {"directory": str(workspace), "outputDir": str(output),
                                  "modelSlxPath": str(workspace / f"{model}.slx"), "modelDir": str(model_dir)}}
            task_path = root / "task.json"
            task_path.write_text(json.dumps(task), encoding="utf-8")
            self.runner.cmd_finish(mock.Mock(task=str(task_path)))
            manifest = json.loads((output / ".tcsd-host" / "execution-manifest.json").read_text(encoding="utf-8"))
            # Real coverage facts, not empty dicts.
            self.assertEqual(manifest["coverage"]["initial"]["models"][model]["condition"]["percent"], 50)
            self.assertEqual(manifest["coverage"]["final"]["models"][model]["mcdc"]["percent"], 66.67)
            # Repair facts derived from validation + proposal.
            self.assertTrue(manifest["coverage"]["repair_required"])
            self.assertTrue(manifest["coverage"]["repair_applied"])
            self.assertEqual(manifest["coverage"]["repair_passes"], 1)
            # MC/DC below threshold with unresolved -> partial, and unresolved recorded.
            self.assertEqual(manifest["completion"], "partial")
            self.assertEqual(manifest["evidence"]["unresolved"][0]["reason_code"], "logic_unreachable")
            # Timeline populated from the 12 checkpoints.
            timeline = json.loads((output / ".tcsd-host" / "timeline.json").read_text(encoding="utf-8"))
            self.assertEqual(len(timeline["events"]), 12)
            # Highest synthesis iteration workbook is delivered to the model dir
            # under the canonical Test0001 name; manifest references the delivered name.
            delivered = model_dir / f"{model}_Test0001_tcsd.xlsx"
            self.assertTrue(delivered.is_file())
            self.assertEqual(delivered.read_bytes(), iter1.read_bytes())
            self.assertEqual(manifest["workbook"], f"outputs/{model}_Test0001_tcsd.xlsx")
            # Artifact manifest lists workbook plus evidence files.
            artifacts = json.loads((output / ".tcsd-host" / "artifact-manifest.json").read_text(encoding="utf-8"))
            roles = [a["role"] for a in artifacts["artifacts"]]
            self.assertIn("workbook", roles)
            self.assertIn("evidence", roles)

    def test_finish_falls_back_to_measured_gaps_when_proposal_has_no_unresolved(self) -> None:
        """ParkCrl B01 regression: 11 uncovered MC/DC vectors stayed out of the
        manifest because the repair proposal recorded no unresolved array; the
        final coverage summary items are the authoritative fallback."""
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workspace = root / "workspace"
            output = workspace / "outputs"
            model_dir = root / "model-dir"
            (output / ".tcsd-host").mkdir(parents=True)
            model_dir.mkdir()
            model = "ParkCrl_A02_B01"
            final_workbook = output / f"{model}_Test0001_tcsd.xlsx"
            wb = Workbook()
            wb.save(final_workbook)
            final = {"models": {model: {
                "model": model, "test_count": 67, "threshold": 80,
                "condition": {"covered": 106, "total": 112, "percent": 94.64, "passed": True},
                "decision": {"covered": 53, "total": 62, "percent": 85.48, "passed": True},
                "mcdc": {"covered": 21, "total": 32, "percent": 65.62, "passed": False},
                "items": [
                    {"coverage_class": "MCDC", "block_path": f"{model}/AND2", "sid": "39",
                     "covered": 1, "total": 3},
                    {"coverage_class": "MCDC", "block_path": f"{model}/OR3", "sid": "75",
                     "covered": 1, "total": 5},
                ],
            }}}
            (output / f"{model}_initial_coverage_summary.json").write_text(
                json.dumps({"models": {model: {"condition": {"percent": 64.29, "passed": False},
                                               "decision": {"percent": 51.61, "passed": False},
                                               "mcdc": {"percent": 34.38, "passed": False}}}}), encoding="utf-8")
            (output / f"{model}_final_coverage_summary.json").write_text(json.dumps(final), encoding="utf-8")
            # Proposal exists but carries no unresolved array.
            (output / f"{model}_agent_coverage_repair_proposal.json").write_text(json.dumps(
                {"schema": "tcsd-agent-coverage-repair-proposal/v1", "jobId": "job-p",
                 "model": model, "tests": [], "unresolved": []}), encoding="utf-8")
            task = {"id": "job-p",
                    "workspace": {"directory": str(workspace), "outputDir": str(output),
                                  "modelSlxPath": str(workspace / f"{model}.slx"), "modelDir": str(model_dir)}}
            task_path = root / "task.json"
            task_path.write_text(json.dumps(task), encoding="utf-8")
            self.runner.cmd_finish(mock.Mock(task=str(task_path)))
            manifest = json.loads((output / ".tcsd-host" / "execution-manifest.json").read_text(encoding="utf-8"))
            unresolved = manifest["evidence"]["unresolved"]
            self.assertEqual(len(unresolved), 2)
            self.assertEqual(unresolved[0]["reason_code"], "measured_uncovered")
            self.assertEqual(unresolved[0]["coverage_class"], "MCDC")
            self.assertEqual(unresolved[0]["block"]["sid"], "39")
            self.assertEqual(manifest["completion"], "partial")

    def test_finish_completion_uses_final_gate_and_merges_measured_gaps(self) -> None:
        """ParkCrl B02 regression: all three final metrics pass (100/100/84.8)
        -> completion=complete even though the INITIAL round was below target;
        proposal-unresolved (logic_unreachable) plus measured gaps merge into
        the evidence without duplication."""
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workspace = root / "workspace"
            output = workspace / "outputs"
            model_dir = root / "model-dir"
            (output / ".tcsd-host").mkdir(parents=True)
            model_dir.mkdir()
            model = "ParkCrl_A02_B02"
            wb = Workbook()
            wb.save(output / f"{model}_Test0001_tcsd.xlsx")
            final = {"models": {model: {
                "model": model, "test_count": 60, "threshold": 80,
                "condition": {"covered": 130, "total": 130, "percent": 100.0, "passed": True},
                "decision": {"covered": 12, "total": 12, "percent": 100.0, "passed": True},
                "mcdc": {"covered": 28, "total": 33, "percent": 84.85, "passed": True},
                "items": [
                    {"coverage_class": "MCDC", "block_path": f"{model}/B02_RPAActv/AND2", "sid": "21",
                     "covered": 2, "total": 3},
                ],
            }}}
            initial = {"models": {model: {"condition": {"percent": 85.38, "passed": True},
                                          "decision": {"percent": 41.67, "passed": False},
                                          "mcdc": {"percent": 27.27, "passed": False}}}}
            (output / f"{model}_initial_coverage_summary.json").write_text(json.dumps(initial), encoding="utf-8")
            (output / f"{model}_final_coverage_summary.json").write_text(json.dumps(final), encoding="utf-8")
            (output / f"{model}_agent_coverage_repair_proposal.json").write_text(json.dumps(
                {"schema": "tcsd-agent-coverage-repair-proposal/v1", "jobId": "job-b2",
                 "model": model, "tests": [], "unresolved": [
                     {"coverage_class": "MCDC", "block": {"path": f"{model}/B02_RPAActv/AND4", "sid": "23"},
                      "reason_code": "logic_unreachable", "evidence": "RPACmd shared root lockout"},
                 ]}), encoding="utf-8")
            task = {"id": "job-b2",
                    "workspace": {"directory": str(workspace), "outputDir": str(output),
                                  "modelSlxPath": str(workspace / f"{model}.slx"), "modelDir": str(model_dir)}}
            task_path = root / "task.json"
            task_path.write_text(json.dumps(task), encoding="utf-8")
            self.runner.cmd_finish(mock.Mock(task=str(task_path)))
            manifest = json.loads((output / ".tcsd-host" / "execution-manifest.json").read_text(encoding="utf-8"))
            # Final gate passed -> complete, initial failure is informational.
            self.assertEqual(manifest["completion"], "complete")
            unresolved = manifest["evidence"]["unresolved"]
            # Proposal entry plus the merged measured gap, no duplication.
            self.assertEqual(len(unresolved), 2)
            self.assertIn("logic_unreachable", [u["reason_code"] for u in unresolved])
            self.assertIn("measured_uncovered", [u["reason_code"] for u in unresolved])
            self.assertEqual(manifest["coverage"]["initial"]["models"][model]["decision"]["percent"], 41.67)


if __name__ == "__main__":
    unittest.main()
